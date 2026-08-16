from datetime import datetime
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QTextCharFormat, QTextCursor
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from PySide6.QtWidgets import (
    QApplication,
    QFileDialog,
    QGroupBox,
    QHBoxLayout,
    QMessageBox,
    QPlainTextEdit,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from config.paths import CODING_VALUE_REPORT_DIR, EXPORT_CODING_FILES_DIR
from core.coding_value import (
    export_coding_value_rows_to_json,
    load_coding_value_rows,
    load_coding_value_rows_from_json,
)
from gui.themes.theme_manager import ThemeManager
from gui.themes.styles.containers.groupbox_style import fn_groupbox_style
from gui.themes.styles.controls.scrollbar_style import fn_apply_scrollbar_style
from gui.utils.text_selection import clear_text_selection_on_focus_out
from gui.widgets.coding_value.coding_value_table import CodingValueTable
from gui.widgets.controls.primary_button import PrimaryButton
from gui.widgets.controls.primary_combobox import PrimaryComboBox
from gui.widgets.controls.primary_lineedit import PrimaryLineEdit


DEFAULT_CODING_DIR = Path("config/Coding")
POSITION_COPY_DATA_PAYLOAD = 3


class CodingValuePanel(QGroupBox):

    def __init__(self):
        super().__init__("Coding Value")

        self._payload_preview_bytes = []
        self._payload_baseline_bytes = []
        self._has_encoded_payload = False
        self._preview_editing = False
        self._setup_ui()
        self._connect_signals()
        self.fn_refresh_theme()

    def _setup_ui(self):
        self.content = QWidget()

        layout = QVBoxLayout(self)
        layout.setContentsMargins(
            12,
            14,
            12,
            12,
        )
        layout.setSpacing(8)

        self.file_path = PrimaryLineEdit(
            placeholder="Select coding Excel file"
        )
        self.btn_browse = PrimaryButton(
            "Browse",
            width=100,
        )
        self.cmb_coding_json = PrimaryComboBox()
        self.cmb_coding_json.setFixedWidth(300)
        self.cmb_coding_json.setMaxVisibleItems(10)
        fn_apply_scrollbar_style(self.cmb_coding_json.view())
        self.cmb_coding_json.setCurrentIndex(-1)
        self.btn_import = PrimaryButton(
            "Import",
            width=100,
        )

        self.file_row = QWidget()
        file_layout = QHBoxLayout(self.file_row)
        file_layout.setContentsMargins(0, 0, 0, 0)
        file_layout.setSpacing(8)
        file_layout.addWidget(self.file_path, 1)
        file_layout.addWidget(self.btn_browse)
        file_layout.addWidget(self.cmb_coding_json)
        file_layout.addWidget(self.btn_import)
        self._refresh_coding_json_options()

        self.txt_coding_value = QPlainTextEdit()
        self.txt_coding_value.setPlaceholderText(
            "Paste coding value payload here"
        )
        self.txt_coding_value.setVerticalScrollBarPolicy(
            Qt.ScrollBarAsNeeded
        )

        self.txt_coding_preview = QPlainTextEdit()
        self.txt_coding_preview.setPlaceholderText(
            "Coding payload preview"
        )
        self.txt_coding_preview.setReadOnly(True)
        self.txt_coding_preview.setVerticalScrollBarPolicy(
            Qt.ScrollBarAsNeeded
        )
        self._set_payload_editor_height()

        self.btn_encode = PrimaryButton(
            "Encode",
            width=100,
        )
        self.btn_copy = PrimaryButton(
            "Copy",
            width=100,
        )
        self.btn_clear = PrimaryButton(
            "Clear",
            width=100,
        )
        self.btn_preview_refresh = PrimaryButton(
            "Refresh",
            width=100,
        )
        self.btn_copy_preview = PrimaryButton(
            "Copy",
            width=100,
        )
        self.btn_preview_edit = PrimaryButton(
            "EDIT",
            width=100,
        )

        self.payload_input_row = QWidget()
        input_layout = QHBoxLayout(self.payload_input_row)
        input_layout.setContentsMargins(0, 0, 0, 0)
        input_layout.setSpacing(8)
        input_layout.addWidget(self.btn_encode, 0, Qt.AlignTop)
        input_layout.addWidget(self.btn_copy, 0, Qt.AlignTop)
        input_layout.addWidget(self.btn_clear, 0, Qt.AlignTop)
        input_layout.addWidget(self.txt_coding_value, 1)

        self.payload_preview_row = QWidget()
        preview_layout = QHBoxLayout(self.payload_preview_row)
        preview_layout.setContentsMargins(0, 0, 0, 0)
        preview_layout.setSpacing(8)
        preview_layout.addWidget(self.btn_preview_refresh, 0, Qt.AlignTop)
        preview_layout.addWidget(self.btn_copy_preview, 0, Qt.AlignTop)
        preview_layout.addWidget(self.btn_preview_edit, 0, Qt.AlignTop)
        preview_layout.addWidget(self.txt_coding_preview, 1)

        self.payload_row = QWidget()
        payload_layout = QVBoxLayout(self.payload_row)
        payload_layout.setContentsMargins(0, 0, 0, 0)
        payload_layout.setSpacing(6)
        payload_layout.addWidget(self.payload_input_row)
        payload_layout.addWidget(self.payload_preview_row)

        self.txt_parameter_filter = PrimaryLineEdit(
            placeholder="Filter parameter"
        )
        filter_action_width = 256
        filter_button_width = (filter_action_width - 8) // 2
        self.btn_filter_no_m = PrimaryButton(
            "No-M",
            width=filter_button_width,
        )
        self.btn_filter_no_m.setFixedWidth(filter_button_width)
        self.btn_refresh_filter = PrimaryButton(
            "Refresh",
            width=filter_button_width,
        )
        self.btn_refresh_filter.setFixedWidth(filter_button_width)

        self.filter_action_row = QWidget()
        self.filter_action_row.setFixedWidth(filter_action_width)
        filter_action_layout = QHBoxLayout(self.filter_action_row)
        filter_action_layout.setContentsMargins(0, 0, 0, 0)
        filter_action_layout.setSpacing(8)
        filter_action_layout.addWidget(self.btn_filter_no_m)
        filter_action_layout.addWidget(self.btn_refresh_filter)

        self.filter_row = QWidget()
        filter_layout = QHBoxLayout(self.filter_row)
        filter_layout.setContentsMargins(0, 0, 0, 0)
        filter_layout.setSpacing(8)
        filter_layout.addWidget(self.txt_parameter_filter, 1)
        filter_layout.addWidget(self.filter_action_row, 0)

        self.table = CodingValueTable()

        self.btn_check = PrimaryButton(
            "CHECK",
            width=110,
        )
        self.btn_copy_data_payload = PrimaryButton(
            "COPY DATA PAYLOAD",
            width=160,
        )
        self.btn_export = PrimaryButton(
            "EXPORT",
            width=110,
        )
        self.btn_table_clear = PrimaryButton(
            "CLEAR",
            width=110,
        )

        self.table_action_panel = QWidget()
        action_layout = QVBoxLayout(self.table_action_panel)
        action_layout.setContentsMargins(0, 0, 0, 0)
        action_layout.setSpacing(8)
        action_layout.addWidget(self.btn_check)
        action_layout.addWidget(self.btn_copy_data_payload)
        action_layout.addWidget(self.btn_export)
        action_layout.addWidget(self.btn_table_clear)
        action_layout.addStretch(1)

        self.txt_working_log = QPlainTextEdit()
        self.txt_working_log.setPlaceholderText(
            "Working log"
        )
        self.txt_working_log.setReadOnly(True)
        self.txt_working_log.setMinimumWidth(220)
        self.txt_working_log.setVerticalScrollBarPolicy(
            Qt.ScrollBarAsNeeded
        )

        self.table_side_panel = QWidget()
        side_layout = QVBoxLayout(self.table_side_panel)
        side_layout.setContentsMargins(0, 0, 0, 0)
        side_layout.setSpacing(8)
        side_layout.addWidget(self.table_action_panel, 0)
        side_layout.addWidget(self.txt_working_log, 1)

        self.table_area = QWidget()
        table_area_layout = QHBoxLayout(self.table_area)
        table_area_layout.setContentsMargins(0, 0, 0, 0)
        table_area_layout.setSpacing(8)
        table_area_layout.addWidget(self.table, 1)
        table_area_layout.addWidget(self.table_side_panel, 0)

        layout.addWidget(self.file_row)
        layout.addWidget(self.payload_row)
        layout.addWidget(self.filter_row)
        layout.addWidget(self.table_area, 1)

        self._install_text_selection_handlers()

    def _install_text_selection_handlers(self):
        clear_text_selection_on_focus_out(
            self,
            [
                self.file_path,
                self.txt_coding_value,
                self.txt_coding_preview,
                self.txt_parameter_filter,
                self.txt_working_log,
            ],
        )

    def _set_payload_editor_height(self):
        self._set_editor_three_line_height(self.txt_coding_value)
        self._set_editor_three_line_height(self.txt_coding_preview)

    @staticmethod
    def _set_editor_three_line_height(editor):
        line_height = editor.fontMetrics().lineSpacing()
        frame_width = editor.frameWidth() * 2
        vertical_padding = 12
        height = line_height * 3 + frame_width + vertical_padding
        editor.setMinimumHeight(height)
        editor.setMaximumHeight(height)

    def _connect_signals(self):
        self.btn_browse.clicked.connect(
            self.browse_file
        )
        self.btn_import.clicked.connect(
            self.import_coding_value
        )
        self.cmb_coding_json.currentIndexChanged.connect(
            self.load_selected_coding_json
        )
        self.btn_encode.clicked.connect(
            self.encode_coding_payload
        )
        self.btn_copy.clicked.connect(
            self.copy_coding_payload
        )
        self.btn_clear.clicked.connect(
            self.clear_coding_payload
        )
        self.btn_preview_refresh.clicked.connect(
            self.refresh_coding_preview
        )
        self.btn_preview_edit.clicked.connect(
            self.toggle_preview_edit_import
        )
        self.btn_copy_preview.clicked.connect(
            self.copy_coding_preview
        )
        self.btn_check.clicked.connect(
            self.check_coding_value
        )
        self.btn_copy_data_payload.clicked.connect(
            self.copy_data_payload
        )
        self.btn_export.clicked.connect(
            self.export_coding_value
        )
        self.btn_table_clear.clicked.connect(
            self.clear_table_coding_values
        )
        self.table.raw_value_changed.connect(
            self.update_payload_preview_from_raw
        )
        self.txt_parameter_filter.textChanged.connect(
            self.filter_parameter_table
        )
        self.btn_filter_no_m.clicked.connect(
            self.filter_no_match_rows
        )
        self.btn_refresh_filter.clicked.connect(
            self.refresh_parameter_filter
        )
        self.txt_coding_value.textChanged.connect(
            self._update_action_states
        )

    def browse_file(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Coding Excel",
            str(DEFAULT_CODING_DIR),
            "Excel Files (*.xlsx *.xlsm *.xltx *.xltm)",
        )

        if file_path:
            self.file_path.setText(file_path)

    def _refresh_coding_json_options(self, selected_path=None):
        selected = str(selected_path) if selected_path else ""
        export_dir = Path(EXPORT_CODING_FILES_DIR)
        json_files = sorted(export_dir.glob("*.json")) if export_dir.exists() else []

        self.cmb_coding_json.blockSignals(True)
        self.cmb_coding_json.clear()
        for json_file in json_files:
            self.cmb_coding_json.addItem(
                json_file.stem,
                str(json_file),
            )

        if selected:
            self.cmb_coding_json.setCurrentIndex(
                self.cmb_coding_json.findData(selected)
            )
        else:
            self.cmb_coding_json.setCurrentIndex(-1)
        self.cmb_coding_json.blockSignals(False)

    def load_selected_coding_json(self):
        json_path = self.cmb_coding_json.currentData()
        if not json_path:
            return

        try:
            rows = load_coding_value_rows_from_json(json_path)
        except Exception as exc:
            QMessageBox.warning(
                self,
                "Coding value",
                f"Cannot import coding JSON file:\n{exc}",
            )
            return

        self._set_coding_rows(rows)

    def _set_coding_rows(self, rows):
        self.table.set_rows(rows)
        self.filter_parameter_table()
        self._update_action_states()

    def import_coding_value(self):
        selected_json = self.cmb_coding_json.currentData()
        excel_path = self.file_path.text().strip()
        if selected_json and not excel_path:
            self.load_selected_coding_json()
            return

        if not excel_path:
            QMessageBox.warning(
                self,
                "Coding value",
                "Please select a coding Excel file before importing.",
            )
            return

        try:
            rows = load_coding_value_rows(
                excel_path,
                "",
            )
        except Exception as exc:
            QMessageBox.warning(
                self,
                "Coding value",
                f"Cannot import coding Excel file:\n{exc}",
            )
            return

        json_path = export_coding_value_rows_to_json(
            excel_path,
            rows,
            EXPORT_CODING_FILES_DIR,
        )
        self._refresh_coding_json_options(json_path)
        self._set_coding_rows(
            load_coding_value_rows_from_json(json_path)
        )
        self.file_path.clear()

    def filter_parameter_table(self):
        self.table.filter_by_parameter(
            self.txt_parameter_filter.text()
        )

    def filter_no_match_rows(self):
        self.txt_parameter_filter.setText("No-M")
        self.filter_parameter_table()

    def refresh_parameter_filter(self):
        self.txt_parameter_filter.clear()
        self.filter_parameter_table()

    def fn_set_crc_value(self, crc_value):
        value = str(crc_value or "").strip()
        if not value:
            return False

        updated = self.table.fn_set_raw_value_by_parameter("crc", value)
        if not updated:
            QMessageBox.warning(
                self,
                "Coding value",
                "Cannot find CRC parameter row.",
            )
            return False

        self._update_action_states()
        return True

    def encode_coding_payload(self):
        self._encode_payload_to_table(
            warn_when_empty=True,
            warn_when_no_rows=True,
        )

    def _encode_payload_to_table(
            self,
            warn_when_empty=False,
            warn_when_no_rows=False,
        ):
        payload = self.txt_coding_value.toPlainText()
        if not payload.strip():
            if warn_when_empty:
                QMessageBox.warning(
                    self,
                    "Coding value",
                    "Please input coding value payload before encoding.",
                )
            self._update_action_states()
            return 0

        updated = self.table.encode_payload(payload)
        if updated > 0:
            self.set_payload_preview(payload)
            self._has_encoded_payload = True
        elif warn_when_no_rows:
            QMessageBox.warning(
                self,
                "Coding value",
                "No raw value row could be updated from this payload.",
            )
        self._update_action_states()
        return updated

    def check_coding_value(self):
        if not self._can_use_table_actions():
            return

        self.table.apply_check_results()
        self._write_check_working_log()
        self._update_action_states()

    def export_coding_value(self):
        if not self._can_use_table_actions():
            return

        output_file = self._choose_coding_report_file()
        if output_file is None:
            return

        try:
            output_file = self._export_coding_report(output_file)
        except PermissionError:
            QMessageBox.warning(
                self,
                "Coding value",
                "Cannot save Coding value report. Please close the Excel file and try again.",
            )
            return

        QMessageBox.information(
            self,
            "Coding value",
            f"Coding value report saved:\n\n{output_file}",
        )

    def _choose_coding_report_file(self):
        default_path = self._default_coding_report_file()
        output_file, _ = QFileDialog.getSaveFileName(
            self,
            "Save Coding Value Report",
            str(default_path),
            "Excel Files (*.xlsx)",
        )
        if not output_file:
            return None

        output_path = Path(output_file)
        if output_path.suffix.lower() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")
        return output_path

    def _default_coding_report_file(self):
        now = datetime.now()
        report_folder = CODING_VALUE_REPORT_DIR / now.strftime("%d-%m-%Y")
        report_folder.mkdir(
            parents=True,
            exist_ok=True,
        )
        return report_folder / (
            f"EEIV_report_Coding value_{now.strftime('%Y%m%d_%H%M%S')}.xlsx"
        )

    def _export_coding_report(self, output_file):
        output_file = Path(output_file)
        output_file.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self._coding_export_sheet_name()
        headers = self._coding_table_headers()
        sheet.append(headers)
        result_index = self._coding_result_column_index(headers)
        for row_index in range(self.table.rowCount()):
            if self.table.isRowHidden(row_index):
                continue

            values = self._coding_table_row_values(row_index)
            sheet.append(values)
            if self._is_no_match_export_row(values, result_index):
                self._apply_no_match_export_fill(sheet[sheet.max_row])

        workbook.save(output_file)
        return output_file

    @staticmethod
    def _coding_result_column_index(headers):
        try:
            return headers.index("Result")
        except ValueError:
            return None

    @staticmethod
    def _is_no_match_export_row(values, result_index):
        return (
            result_index is not None
            and result_index < len(values)
            and values[result_index] == "No-M"
        )

    @staticmethod
    def _apply_no_match_export_fill(row_cells):
        warning = ThemeManager.fn_colors().WARNING.replace("#", "")
        fill = PatternFill(
            fill_type="solid",
            fgColor=warning,
        )
        for cell in row_cells:
            cell.fill = fill

    def _coding_export_sheet_name(self):
        return self._safe_excel_sheet_name(
            f"Coding value_{self._coding_export_identifier()}"
        )

    def _coding_export_identifier(self):
        payload_bytes = self._payload_baseline_bytes or self._payload_preview_bytes
        return " ".join(
            f"{value:02X}"
            for value in payload_bytes[:3]
        )

    def _coding_table_headers(self):
        return [
            self.table.horizontalHeaderItem(column).text()
            for column in range(self.table.columnCount())
        ]

    def _coding_table_row_values(self, row_index):
        values = []
        for column in range(self.table.columnCount()):
            widget = self.table.cellWidget(row_index, column)
            if hasattr(widget, "currentText"):
                values.append(widget.currentText())
                continue

            item = self.table.item(row_index, column)
            values.append(item.text() if item is not None else "")

        return values

    @staticmethod
    def _safe_excel_sheet_name(name):
        invalid_chars = r'[]:*?/\\'
        for char in invalid_chars:
            name = name.replace(char, "_")

        return name[:31] or "Coding value"

    def copy_coding_payload(self):
        QApplication.clipboard().setText(
            self.txt_coding_value.toPlainText()
        )

    def copy_coding_preview(self):
        preview_text = self.txt_coding_preview.toPlainText()
        if not preview_text.strip():
            return

        QApplication.clipboard().setText(preview_text)

    def copy_data_payload(self):
        preview_text = self.txt_coding_preview.toPlainText()
        if not preview_text.strip():
            return

        try:
            payload_bytes = self.table._parse_payload_bytes(preview_text)
        except ValueError:
            return

        data_payload = payload_bytes[POSITION_COPY_DATA_PAYLOAD:]
        if not data_payload:
            return

        QApplication.clipboard().setText(
            self._format_payload_bytes(data_payload)
        )

    def clear_coding_payload(self):
        self.txt_coding_value.clear()
        self.clear_coding_preview()

    def clear_table_coding_values(self):
        self.table.clear_rows()
        self.txt_working_log.clear()
        self.clear_coding_preview()

    def refresh_coding_preview(self):
        if self._encode_payload_to_table() == 0:
            self.set_payload_preview(
                self.txt_coding_value.toPlainText()
            )

    def toggle_preview_edit_import(self):
        if self._preview_editing:
            self.import_preview_payload()
            return

        self.start_preview_edit()

    def start_preview_edit(self):
        if not self._can_edit_preview_payload():
            return

        self._preview_editing = True
        self.txt_coding_preview.setReadOnly(False)
        self.btn_preview_edit.setText("IMPORT")
        self.txt_coding_preview.setFocus()
        self._update_action_states()

    def import_preview_payload(self):
        payload = self.txt_coding_preview.toPlainText()
        if not payload.strip():
            return

        updated = self.table.encode_payload(
            payload,
            update_original=False,
        )
        if updated <= 0:
            QMessageBox.warning(
                self,
                "Coding value",
                "No raw value row could be updated from this payload.",
            )
            return

        self._payload_preview_bytes = self.table._parse_payload_bytes(payload)
        self.txt_coding_preview.setPlainText(
            self._format_payload_bytes(self._payload_preview_bytes)
        )
        self.txt_coding_preview.setExtraSelections([])
        if self._has_check_results():
            self._write_check_working_log()
        self._finish_preview_edit()

    def _finish_preview_edit(self):
        self._preview_editing = False
        self.txt_coding_preview.setReadOnly(True)
        self.btn_preview_edit.setText("EDIT")
        self._update_action_states()

    def _can_edit_preview_payload(self):
        return (
            self._can_use_table_actions()
            and self._has_check_results()
            and bool(self.txt_coding_preview.toPlainText().strip())
        )

    def _has_check_results(self):
        return self.table.columnCount() >= 9

    def clear_coding_preview(self):
        self.txt_coding_preview.clear()
        self.txt_coding_preview.setReadOnly(True)
        self.btn_preview_edit.setText("EDIT")
        self._preview_editing = False
        self.txt_coding_preview.setExtraSelections([])
        self._payload_preview_bytes = []
        self._payload_baseline_bytes = []
        self._has_encoded_payload = False
        self._update_action_states()

    def _update_action_states(self):
        has_input_payload = bool(self.txt_coding_value.toPlainText().strip())
        has_preview_payload = bool(self.txt_coding_preview.toPlainText().strip())
        has_table_rows = self.table.rowCount() > 0
        can_encode = has_table_rows and has_input_payload
        can_use_preview = has_preview_payload
        can_use_table_actions = self._can_use_table_actions()

        self.btn_encode.setEnabled(can_encode)
        self.btn_copy.setEnabled(has_input_payload)
        self.btn_clear.setEnabled(has_input_payload)
        self.btn_preview_refresh.setEnabled(can_encode)
        self.btn_copy_preview.setEnabled(can_use_preview)
        self.btn_preview_edit.setEnabled(self._can_edit_preview_payload())
        self.btn_filter_no_m.setEnabled(
            can_use_table_actions and self._has_check_results()
        )
        self.btn_refresh_filter.setEnabled(has_table_rows)
        self.btn_check.setEnabled(can_use_table_actions)
        self.btn_copy_data_payload.setEnabled(
            len(self._payload_preview_bytes) > 3
        )
        self.btn_export.setEnabled(can_use_table_actions)
        self.btn_table_clear.setEnabled(can_use_table_actions)

    def _can_use_table_actions(self):
        return (
            self._has_encoded_payload
            and self.table.rowCount() > 0
            and bool(self._payload_preview_bytes)
        )

    def _write_check_working_log(self):
        self.txt_working_log.setPlainText(
            self._build_check_working_log()
        )

    def _build_check_working_log(self):
        changes = []
        for row_index in range(self.table.rowCount()):
            result_item = self.table.item(row_index, 8)
            if result_item is None or result_item.text() != "No-M":
                continue

            byte_item = self.table.item(row_index, 1)
            before_decoded_item = self.table.item(row_index, 7)
            current_raw_item = self.table.item(row_index, 4)
            current_raw = current_raw_item.text() if current_raw_item else ""
            byte_text = byte_item.text() if byte_item else str(row_index)
            before_decoded = (
                before_decoded_item.text()
                if before_decoded_item is not None
                else ""
            )
            after_decoded = self.table._decoded_text_for_raw(
                row_index,
                current_raw,
            )
            changes.append(
                f"Byte {byte_text}: {before_decoded} -> {after_decoded}"
            )

        lines = [
            f"Total No-M: {len(changes)}",
            "",
            "Changes: Decode value before -> after",
            "",
        ]
        lines.extend(changes)
        lines.extend([
            "",
            (
                "CRC: "
                f"{self._crc_text(self._payload_baseline_bytes)} -> "
                f"{self._crc_text(self._payload_preview_bytes)}"
            ),
        ])
        return "\n".join(lines)

    @staticmethod
    def _crc_text(payload_bytes):
        if not payload_bytes:
            return ""

        return f"{payload_bytes[-1]:02X}"

    def set_payload_preview(self, payload):
        self._payload_preview_bytes = self.table._parse_payload_bytes(payload)
        self._payload_baseline_bytes = list(self._payload_preview_bytes)
        self.txt_coding_preview.setPlainText(
            self._format_payload_bytes(self._payload_preview_bytes)
        )
        self.txt_coding_preview.setExtraSelections([])

    def update_payload_preview_from_raw(
            self,
            _row_index,
            byte_pos,
            bit_pos,
            bit_length,
            raw_value,
        ):
        changed_bytes = self._apply_raw_to_payload_bytes(
            byte_pos,
            bit_pos,
            bit_length,
            raw_value,
        )
        if not changed_bytes:
            return

        self.txt_coding_preview.setPlainText(
            self._format_payload_bytes(self._payload_preview_bytes)
        )
        self._highlight_payload_bytes(changed_bytes)

    def _apply_raw_to_payload_bytes(
            self,
            byte_pos,
            bit_pos,
            bit_length,
            raw_value,
        ):
        if byte_pos < 0 or byte_pos >= len(self._payload_preview_bytes):
            return []

        tokens = self.table._hex_tokens(raw_value)
        if not tokens:
            return []

        raw_int = int("".join(tokens), 16)
        if bit_pos == 0 and bit_length % 8 == 0:
            byte_count = bit_length // 8
            if byte_count <= 0 or byte_pos + byte_count > len(self._payload_preview_bytes):
                return []
            raw_bytes = raw_int.to_bytes(byte_count, byteorder="big")
            for offset, value in enumerate(raw_bytes):
                self._payload_preview_bytes[byte_pos + offset] = value
            return list(range(byte_pos, byte_pos + byte_count))

        byte_count = (bit_pos + bit_length + 7) // 8
        if byte_count <= 0 or byte_pos + byte_count > len(self._payload_preview_bytes):
            return []

        selected = self._payload_preview_bytes[byte_pos:byte_pos + byte_count]
        container = int.from_bytes(bytes(selected), byteorder="little")
        mask = ((1 << bit_length) - 1) << bit_pos
        container = (container & ~mask) | ((raw_int << bit_pos) & mask)
        merged = container.to_bytes(byte_count, byteorder="little")
        for offset, value in enumerate(merged):
            self._payload_preview_bytes[byte_pos + offset] = value
        return list(range(byte_pos, byte_pos + byte_count))

    @staticmethod
    def _format_payload_bytes(payload_bytes):
        return " ".join(
            f"{value:02X}"
            for value in payload_bytes
        )

    def _highlight_payload_bytes(self, byte_indexes):
        selections = []
        color = QColor("#FFE08A")
        for byte_index in byte_indexes:
            start = byte_index * 3
            cursor = QTextCursor(self.txt_coding_preview.document())
            cursor.setPosition(start)
            cursor.movePosition(
                QTextCursor.Right,
                QTextCursor.KeepAnchor,
                2,
            )
            selection = QTextEdit.ExtraSelection()
            selection.cursor = cursor
            selection.format = QTextCharFormat()
            selection.format.setBackground(color)
            selections.append(selection)

        self.txt_coding_preview.setExtraSelections(selections)

    def fn_refresh_theme(self):
        self.setStyleSheet(
            fn_groupbox_style()
        )
        self.file_path.fn_refresh_theme()
        self.txt_parameter_filter.fn_refresh_theme()
        self.btn_filter_no_m.fn_refresh_theme()
        self.btn_refresh_filter.fn_refresh_theme()
        self.btn_browse.fn_refresh_theme()
        self.btn_import.fn_refresh_theme()
        self.btn_encode.fn_refresh_theme()
        self.btn_copy.fn_refresh_theme()
        self.btn_clear.fn_refresh_theme()
        self.btn_preview_refresh.fn_refresh_theme()
        self.btn_copy_preview.fn_refresh_theme()
        self.btn_preview_edit.fn_refresh_theme()
        self.btn_check.fn_refresh_theme()
        self.btn_copy_data_payload.fn_refresh_theme()
        self.btn_export.fn_refresh_theme()
        self.btn_table_clear.fn_refresh_theme()
        self.table.fn_refresh_theme()
        self._update_action_states()
        fn_apply_scrollbar_style(
            self.txt_coding_value
        )
        fn_apply_scrollbar_style(
            self.txt_coding_preview
        )
        fn_apply_scrollbar_style(
            self.txt_working_log
        )
