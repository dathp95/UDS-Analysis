"""Read-only offline reference extraction. Reprs/summaries never expose seed or key."""
from dataclasses import dataclass, field
from math import isfinite
from collections import Counter


@dataclass(frozen=True)
class SecurityReferenceVector:
    ecu: str
    session: int
    subfunction: int
    seed_time: float
    key_time: float
    seed: bytes = field(repr=False)
    key: bytes = field(repr=False)


@dataclass(frozen=True)
class ReferenceAnalysis:
    vectors: tuple[SecurityReferenceVector, ...] = field(repr=False)
    classifications: dict

    def summary(self):
        lengths = Counter((v.ecu, v.subfunction, len(v.seed), len(v.key)) for v in self.vectors)
        return dict(exchanges=len(self.vectors), classifications=self.classifications,
                    lengths=[dict(ecu=e, subfunction=s, seed_bytes=a, key_bytes=b, count=n)
                             for (e, s, a, b), n in sorted(lengths.items())])


def _hex(value):
    if not isinstance(value, str) or len(value) > 12285:
        return b''
    try:
        return bytes.fromhex(value)
    except ValueError:
        return b''


def extract_reference_rows(rows):
    """Pair complete same-ECU exchanges in source order, never by payload similarity.

    Any intervening same-ECU request consumes/invalidates a candidate. Session
    changes and backwards timestamps cannot link two otherwise similar vectors.
    Response completion (request Time + RT) must precede key request Time.
    """
    columns = None
    pending, sessions, previous = {}, Counter(), {}
    counts, vectors = Counter(), []
    for row in rows:
        if columns is None:
            names = {str(value).strip(): i for i, value in enumerate(row) if value is not None}
            if {'ECU', 'Time', 'Request', 'Response', 'RT (ms)'} <= names.keys():
                columns = names
            continue
        def cell(name):
            index = columns[name]
            return row[index] if index < len(row) else None
        ecu, timestamp, rt = cell('ECU'), cell('Time'), cell('RT (ms)')
        if not isinstance(ecu, str) or not ecu.strip():
            continue
        ecu = ecu.strip()
        candidate = pending.pop(ecu, None)
        if candidate is not None:
            counts['unmatched_seed'] += 1
        if (type(timestamp) not in (int, float) or not isfinite(timestamp)
                or type(rt) not in (int, float) or not isfinite(rt) or rt < 0):
            counts['invalid_time'] += 1
            continue
        if timestamp < previous.get(ecu, -float('inf')):
            sessions[ecu] += 1
            counts['backwards_time'] += 1
            candidate = None
        previous[ecu] = timestamp
        request, response = _hex(cell('Request')), _hex(cell('Response'))
        if request[:1] == b'\x10':
            sessions[ecu] += 1
        if request[:1] != b'\x27':
            continue
        if response[:2] == b'\x7f\x27':
            counts[('pending' if response[2] == 0x78 else 'negative') if len(response) == 3 else 'malformed_negative'] += 1
        elif response[:1] == b'\x67':
            counts['positive'] += 1
        else:
            counts['missing_or_invalid_response'] += 1
        if request in (b'\x27\x01', b'\x27\x03'):
            subfunction = request[1]
            if response[:2] == bytes((0x67, subfunction)) and 3 <= len(response) <= 4095:
                pending[ecu] = (subfunction, response[2:], timestamp, timestamp + rt/1000, sessions[ecu])
            else:
                counts['invalid_seed'] += 1
        elif len(request) > 2 and request[1] in (2, 4):
            if (candidate is not None and candidate[0]+1 == request[1]
                    and response == bytes((0x67, request[1])) and timestamp >= candidate[3]
                    and candidate[4] == sessions[ecu]):
                vectors.append(SecurityReferenceVector(ecu, sessions[ecu], candidate[0],
                                                       candidate[2], timestamp, candidate[1], request[2:]))
                counts['unmatched_seed'] -= 1
            else:
                counts['unmatched_key'] += 1
    if columns is None:
        raise ValueError('Expected ALL_ECUS transaction headers were not found.')
    counts['unmatched_seed'] += len(pending)
    return ReferenceAnalysis(tuple(vectors), dict(counts))


def read_reference_workbook(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if 'ALL_ECUS' not in workbook.sheetnames:
            raise ValueError('Reference workbook has no ALL_ECUS sheet.')
        return extract_reference_rows(workbook['ALL_ECUS'].values)
    finally:
        workbook.close()
