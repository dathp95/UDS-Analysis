import tempfile
import unittest
from pathlib import Path

import can
from openpyxl import load_workbook

from core.asc_reader import get_log_channels, load_log, reassemble_isotp
from core.pipeline import run_pipeline
from core.report_export import export_workbook
from core.transaction_formatter import fn_build_transaction_rows
from models.ecu import ECU
from models.vehicle import Vehicle
from tests.core.test_dsa_txt_parser import HEADER, dsa_frame


class DSAPipelineTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[2])
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "diagnostic.txt"
        self.vehicle = Vehicle(name="Test", ecus=[ECU(name="ECU", request_id="682", response_id="602")])
        self.records = [
            ("00:00:20.925.490", "682", "02 10 03 AA AA AA AA AA"),
            ("00:00:20.929.897", "602", "06 50 03 00 32 00 C8 CC"),
            ("00:00:22.915.223", "682", "03 22 F1 68 AA AA AA AA"),
            ("00:00:22.919.897", "602", "03 7F 22 78 CC CC CC CC"),
            ("00:00:22.924.198", "602", "04 62 F1 68 00 CC CC CC"),
            ("00:00:23.000.001", "682", "02 27 03 AA AA AA AA AA"),
            ("00:00:23.001.002", "602", "10 12 67 03 AA 6F BC AD"),
            ("00:00:23.001.010", "682", "30 00 00 AA AA AA AA AA"),
            ("00:00:23.002.003", "602", "21 F1 77 23 21 BB C3 61"),
            ("00:00:23.003.004", "602", "22 0F CA C7 49 BD CC CC"),
        ]
        self.path.write_text(HEADER + "".join(
            dsa_frame(timestamp=t, can_id=cid, data=data) for t, cid, data in self.records
        ), encoding="utf-8")

    def test_session_did_nrc_pending_and_microsecond_response_time_use_existing_pipeline(self):
        result = run_pipeline(str(self.path), self.vehicle, channel=0)
        session, did, security = result["transactions"]
        self.assertEqual(session["request"]["payload"], "10 03")
        self.assertEqual(session["positive_response"]["payload"], "50 03 00 32 00 C8")
        self.assertAlmostEqual(session["response_time"], 0.004407, places=9)
        self.assertEqual(did["request"]["payload"], "22 F1 68")
        self.assertEqual(did["negative_responses"][0]["payload"], "7F 22 78")
        self.assertEqual(did["response_pending_count"], 1)
        self.assertEqual(did["positive_response"]["payload"], "62 F1 68 00")
        self.assertEqual(did["status"], "OK")
        self.assertAlmostEqual(did["response_time"], 0.008975, places=9)
        self.assertEqual(security["request"]["payload"], "27 03")
        self.assertEqual(security["positive_response"]["payload"],
                         "67 03 AA 6F BC AD F1 77 23 21 BB C3 61 0F CA C7 49 BD")
        self.assertAlmostEqual(security["response_time"], 0.001001, places=9)

    def test_multiframe_reassembly_is_isolated_by_selected_channel_and_can_id(self):
        with self.path.open("a", encoding="utf-8") as log:
            log.write(dsa_frame(channel=1, can_id="602", data="10 12 67 03 11 11 11 11"))
            log.write(dsa_frame(channel=1, can_id="602", data="21 11 11 11 11 11 11 11"))
            log.write(dsa_frame(can_id="603", data="21 22 22 22 22 22 22 22"))
        result = run_pipeline(str(self.path), self.vehicle, channel=0)
        self.assertEqual(len(result["transactions"]), 3)
        self.assertIn("AA 6F BC AD", result["transactions"][-1]["positive_response"]["payload"])
        self.assertEqual(result["transactions"][-1]["positive_response"]["can_id"], 0x602)
        self.assertEqual(reassemble_isotp(load_log(self.path, channel=1)), [])

    def test_valid_zero_dlc_and_incomplete_first_frame_do_not_crash_uds(self):
        with self.path.open("a", encoding="utf-8") as log:
            log.write(dsa_frame(dlc=0, data=""))
            log.write(dsa_frame(dlc=1, data="10"))
        self.assertEqual(len(run_pipeline(str(self.path), self.vehicle, channel=0)["transactions"]), 3)

    def test_txt_result_uses_existing_row_formatter_and_excel_export(self):
        result = run_pipeline(str(self.path), self.vehicle, channel=0)
        rows = fn_build_transaction_rows(result["transactions"])
        self.assertEqual(rows[1]["Request"], "22 F1 68")
        self.assertEqual(rows[1]["Response"], "62 F1 68 00")
        self.assertEqual(rows[1]["RT (ms)"], 8.97)
        output = self.path.with_suffix(".xlsx")
        self.assertTrue(export_workbook(result["summary"], result["ecu_reports"], result["transactions"], output))
        workbook = load_workbook(output)
        self.addCleanup(workbook.close)
        self.assertIn("ALL_ECUS", workbook.sheetnames)
        self.assertEqual(workbook["ALL_ECUS"]["D4"].value, "22 F1 68")

    def test_blf_and_asc_real_files_continue_through_the_same_pipeline(self):
        blf = self.path.with_suffix(".blf")
        with can.BLFWriter(str(blf)) as writer:
            for index, (_, can_id, data) in enumerate(self.records[:2]):
                writer.on_message_received(can.Message(
                    timestamp=1000.0 + index * 0.01, arbitration_id=int(can_id, 16),
                    is_extended_id=False, channel=0, data=bytes.fromhex(data),
                ))
        self.assertEqual(get_log_channels(blf), [1])
        result = run_pipeline(str(blf), self.vehicle, channel=1)
        self.assertEqual(result["transactions"][0]["request"]["payload"], "10 03")
        asc = blf.with_suffix(".asc")
        self.assertTrue(asc.is_file())
        self.assertEqual(get_log_channels(asc), [1])
        self.assertEqual(run_pipeline(str(asc), self.vehicle, channel=1), result)


if __name__ == "__main__":
    unittest.main()
