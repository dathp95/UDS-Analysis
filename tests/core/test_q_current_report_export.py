import base64
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from core.excel_styles import (
    COLOR_FAIL,
    COLOR_HEADER,
    COLOR_PASS,
    COLOR_SECTION,
)
from core.q_current_analysis import QCurrentAnalysisResult, QCurrentWakeUpInterval
from core.q_current_report_export import (
    QCurrentReportData,
    QCurrentReportSettings,
    export_q_current_report,
)
from core.sleep_current_database import CurrentSample


_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


def _cell_color(cell):
    value = cell.fill.fgColor.rgb
    return "" if value is None else str(value).upper()


def _assert_color_endswith(testcase, cell, color):
    testcase.assertTrue(
        _cell_color(cell).endswith(color),
        f"Expected {cell.coordinate} color to end with {color}, got {_cell_color(cell)}",
    )


class QCurrentReportExportTests(unittest.TestCase):

    def test_exports_data_review_and_report_workbook(self):
        samples = [
            CurrentSample(1, 1, "26-09-11 11:41:29.036", -0.0169, -16.9),
            CurrentSample(2, 1, "26-09-11 11:41:30.079", -0.0173, -17.3),
            CurrentSample(3, 1, "26-09-11 11:41:31.101", -0.3150, -315.0),
        ]
        analysis_result = QCurrentAnalysisResult(
            result_status="PASSED",
            average_sleep_current_ma=17.1,
            minimum_sleep_current_ma=16.9,
            maximum_sleep_current_ma=17.3,
            wake_up_event_count=1,
            total_wake_up_duration_s=2.0,
            average_wake_up_duration_s=2.0,
            maximum_wake_up_duration_s=2.0,
            average_wake_up_interval_s=None,
            minimum_wake_up_interval_s=None,
            maximum_wake_up_interval_s=None,
            longest_continuous_sleep_s=1.043,
            start_time_s=0.0,
            end_time_s=2.065,
            source_start_time="26-09-11 11:41:29.036",
            source_end_time="26-09-11 11:41:31.101",
            duration_s=2.065,
            total_samples=3,
            sample_interval_s=1.033,
            wake_up_intervals=(QCurrentWakeUpInterval(1.0, 3.0, 2.0, 315.0),),
        )
        report_data = QCurrentReportData(
            dataset_name="VF8_Sleep_Current",
            samples=samples,
            analysis_result=analysis_result,
            settings=QCurrentReportSettings(
                standard_current_ma=30.0,
                wake_up_limit_ma=300.0,
                wake_duration_s=2.0,
            ),
            chart_png=_TINY_PNG,
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            output_file = Path(tmpdir) / "report.xlsx"
            self.assertTrue(export_q_current_report(report_data, output_file))

            workbook = load_workbook(output_file)
            self.assertEqual(workbook.sheetnames, ["Data Review", "Report"])
            self.assertEqual(workbook.properties.creator, "Dat Tran")
            self.assertEqual(workbook.properties.title, "Q Current Analysis Report")
            self.assertEqual(workbook.properties.subject, "Vehicle Sleep Current Analysis")

            data_sheet = workbook["Data Review"]
            self.assertEqual(data_sheet["A1"].value, "Time")
            self.assertEqual(data_sheet["B1"].value, "Current (mA)")
            self.assertEqual(data_sheet.freeze_panes, "A2")
            self.assertEqual(data_sheet.auto_filter.ref, "A1:B4")
            self.assertEqual(data_sheet["A2"].value, "26-09-11 11:41:29.036")
            self.assertEqual(data_sheet["B2"].value, -16.9)
            self.assertEqual(data_sheet["B2"].number_format, "0.00")
            self.assertTrue(data_sheet["A1"].font.bold)
            _assert_color_endswith(self, data_sheet["A1"], COLOR_HEADER)
            self.assertEqual(data_sheet["A1"].alignment.horizontal, "center")
            self.assertEqual(data_sheet["A1"].border.left.style, "thin")

            report_sheet = workbook["Report"]
            self.assertEqual(report_sheet["A1"].value, "Q CURRENT ANALYSIS REPORT")
            self.assertEqual(report_sheet["A2"].value, "by Dat Tran")
            self.assertEqual(report_sheet["A3"].value, "Dataset: VF8_Sleep_Current")
            values = [cell.value for row in report_sheet.iter_rows() for cell in row]
            self.assertIn("TEST SETTINGS", values)
            self.assertIn("RESULT", values)
            self.assertIn("CURRENT STATISTICS", values)
            self.assertIn("WAKE-UP STATISTICS", values)
            self.assertIn("TEST INFORMATION", values)
            self.assertIn("CURRENT CHART", values)
            self.assertIn("PASSED", values)
            section_cell = next(
                cell
                for row in report_sheet.iter_rows()
                for cell in row
                if cell.value == "TEST SETTINGS"
            )
            result_cell = next(
                cell
                for row in report_sheet.iter_rows()
                for cell in row
                if cell.value == "PASSED"
            )
            _assert_color_endswith(self, section_cell, COLOR_SECTION)
            self.assertTrue(section_cell.font.bold)
            self.assertEqual(section_cell.font.color.rgb[-6:].upper(), "FFFFFF")
            _assert_color_endswith(self, result_cell, COLOR_PASS)
            self.assertIn(30.0, values)
            self.assertIn(300.0, values)
            self.assertIn(2.0, values)
            self.assertIn("11:41:29", values)
            self.assertIn("11:41:31", values)
            self.assertEqual(len(report_sheet._images), 1)

    def test_unavailable_values_are_written_as_na(self):
        report_data = QCurrentReportData(
            dataset_name="Dataset",
            samples=[],
            analysis_result=QCurrentAnalysisResult(
                result_status="FAILED",
                average_sleep_current_ma=None,
                minimum_sleep_current_ma=None,
                maximum_sleep_current_ma=None,
                wake_up_event_count=0,
                total_wake_up_duration_s=0.0,
                average_wake_up_duration_s=None,
                maximum_wake_up_duration_s=None,
                average_wake_up_interval_s=None,
                minimum_wake_up_interval_s=None,
                maximum_wake_up_interval_s=None,
                longest_continuous_sleep_s=None,
                start_time_s=None,
                end_time_s=None,
                source_start_time=None,
                source_end_time=None,
                duration_s=None,
                total_samples=0,
                sample_interval_s=None,
                wake_up_intervals=(),
            ),
            settings=QCurrentReportSettings(30.0, 300.0, 2.0),
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            output_file = Path(tmpdir) / "empty.xlsx"
            self.assertTrue(export_q_current_report(report_data, output_file))
            workbook = load_workbook(output_file)
            values = [cell.value for row in workbook["Report"].iter_rows() for cell in row]
            self.assertIn("N/A", values)
            self.assertEqual(workbook["Data Review"].auto_filter.ref, "A1:B1")
            values = [cell for row in workbook["Report"].iter_rows() for cell in row]
            failed_cell = next(cell for cell in values if cell.value == "FAILED")
            _assert_color_endswith(self, failed_cell, COLOR_FAIL)


if __name__ == "__main__":
    unittest.main()