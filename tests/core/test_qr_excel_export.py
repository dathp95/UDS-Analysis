from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image

from core.qr_excel_export import export_qr_codes_to_excel, parse_qr_codes
from core.qr_generator import generate_qr_png_bytes


class QRExcelExportTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[2])
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.path = self.root / "codes.xlsx"

    def test_parse_ignores_blank_lines_and_preserves_spaces_duplicates_and_zeros(self):
        self.assertEqual(parse_qr_codes(" \n 00001 \n a  b\n00001\n\n"), ["00001", "a  b", "00001"])

    def test_empty_and_over_limit_rejected_before_any_image_generation(self):
        for values in ([], ["", " "], ["QR"] * 51):
            with self.subTest(count=len(values)), patch("core.qr_excel_export.generate_qr_png_bytes") as generate:
                with self.assertRaises(ValueError):
                    export_qr_codes_to_excel(values, self.path)
                generate.assert_not_called()
                self.assertFalse(self.path.exists())
                self.assertEqual(list(self.root.iterdir()), [])

    def test_sheet_boundaries_images_order_and_print_settings(self):
        for count in (1, 3, 24, 25, 48, 50):
            with self.subTest(count=count):
                values = [f"QR-{i:03d}" for i in range(count)]
                result = export_qr_codes_to_excel(values, self.path)
                self.assertEqual(result.code_count, count)
                self.assertEqual(result.sheet_count, (count + 23) // 24)
                with ZipFile(self.path) as archive:
                    self.assertIsNone(archive.testzip())
                    self.assertEqual(len([name for name in archive.namelist() if name.startswith("xl/media/")]), count)
                workbook = load_workbook(self.path)
                try:
                    self.assertEqual(workbook.sheetnames, [f"QR_{i + 1:03d}" for i in range(result.sheet_count)])
                    actual = []
                    for sheet in workbook:
                        self.assertGreater(len(sheet._images), 0)
                        self.assertLessEqual(len(sheet._images), 24)
                        self.assertFalse(sheet.sheet_view.showGridLines)
                        self.assertEqual(sheet.page_setup.orientation, "portrait")
                        self.assertEqual(sheet.page_setup.paperSize, int(sheet.PAPERSIZE_A4))
                        self.assertEqual((sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight), (1, 1))
                        for i, image in enumerate(sheet._images):
                            row, column = 2 * (i // 3), i % 3
                            self.assertEqual((image.anchor._from.row, image.anchor._from.col), (row, column))
                            self.assertEqual(image.anchor.ext.cx, 135 * 9525)
                            self.assertGreater(image.anchor._from.colOff, 0)
                            label = sheet.cell(row + 2, column + 1)
                            actual.append(label.value)
                            self.assertEqual(label.data_type, "s")
                    self.assertEqual(actual, values)
                finally:
                    workbook.close()

    def test_labels_are_strings_and_png_payloads_match_each_trimmed_value(self):
        values = [" 000123 ", "A  B", "\u0110i\u1ec7n \u4e2d", "=1+1", "+123", "-123", "@value", "#N/A", "000123"]
        expected = [value.strip() for value in values]
        export_qr_codes_to_excel(values, self.path)
        workbook = load_workbook(self.path, data_only=False)
        try:
            for i, value in enumerate(expected):
                sheet = workbook.active
                label = sheet.cell(2 * (i // 3) + 2, i % 3 + 1)
                self.assertEqual(label.value, value)
                self.assertEqual(label.data_type, "s")
                self.assertEqual(label.number_format, "@")
                with Image.open(BytesIO(sheet._images[i]._data())) as actual, \
                     Image.open(BytesIO(generate_qr_png_bytes(value))) as original:
                    self.assertEqual(actual.size, original.size)
                    self.assertEqual(actual.tobytes(), original.tobytes())
            with ZipFile(self.path) as archive:
                self.assertNotIn(b"<f>", archive.read("xl/worksheets/sheet1.xml"))
        finally:
            workbook.close()

    def test_invalid_path_does_not_generate_images(self):
        with patch("core.qr_excel_export.generate_qr_png_bytes") as generate:
            with self.assertRaises((ValueError, OSError)):
                export_qr_codes_to_excel(["QR"], self.root / "missing" / "codes.xlsx")
            generate.assert_not_called()

    def test_generation_save_and_replace_failures_preserve_destination_and_clean_temp(self):
        original = b"existing destination"
        for operation in ("generate_qr_png_bytes", "Workbook.save", "os.replace"):
            with self.subTest(operation=operation):
                self.path.write_bytes(original)
                with patch("core.qr_excel_export." + operation, side_effect=OSError("simulated failure")):
                    with self.assertRaises(OSError):
                        export_qr_codes_to_excel(["QR"], self.path)
                self.assertEqual(self.path.read_bytes(), original)
                self.assertEqual(list(self.root.iterdir()), [self.path])

    def test_success_atomically_replaces_existing_file(self):
        self.path.write_bytes(b"old file")
        export_qr_codes_to_excel(["00001"], self.path)
        workbook = load_workbook(self.path)
        try:
            self.assertEqual(workbook.active["A2"].value, "00001")
        finally:
            workbook.close()
        self.assertEqual(list(self.root.iterdir()), [self.path])

    def test_partial_save_failure_leaves_no_final_or_temporary_file(self):
        def partial_save(stream):
            stream.write(b"partial workbook")
            raise OSError("disk full")

        with patch("core.qr_excel_export.Workbook.save", side_effect=partial_save):
            with self.assertRaisesRegex(OSError, "disk full"):
                export_qr_codes_to_excel(["QR"], self.path)
        self.assertEqual(list(self.root.iterdir()), [])

    def test_image_buffers_are_closed_after_failure(self):
        buffers = []

        def buffer(data):
            stream = BytesIO(data)
            buffers.append(stream)
            return stream

        with patch("core.qr_excel_export.BytesIO", side_effect=buffer), patch(
            "core.qr_excel_export.Workbook.save", side_effect=OSError("save failed")
        ):
            with self.assertRaises(OSError):
                export_qr_codes_to_excel(["QR1", "QR2"], self.path)
        self.assertEqual(len(buffers), 2)
        self.assertTrue(all(stream.closed for stream in buffers))

    def test_invalid_xml_input_is_rejected_before_generation(self):
        with patch("core.qr_excel_export.generate_qr_png_bytes") as generate:
            with self.assertRaises(ValueError):
                export_qr_codes_to_excel(["valid", "invalid\x00text"], self.path)
            generate.assert_not_called()
