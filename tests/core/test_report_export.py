import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from core.excel_styles import COLOR_HEADER
from core.report_export import export_workbook


def _cell_color(cell):
    value = cell.fill.fgColor.rgb
    return "" if value is None else str(value).upper()


class ReportExportTests(unittest.TestCase):

    def test_export_workbook_applies_shared_table_style_to_all_sheets(self):
        summary = [
            {
                "ecu": "BCM",
                "total": 1,
                "pass": 1,
            }
        ]
        transactions = [
            {
                "ecu": "BCM",
                "display_name": "Read VIN",
                "request": {
                    "timestamp": 0.123,
                    "payload": "22 F1 90",
                },
                "positive_response": {
                    "payload": "62 F1 90",
                },
                "response_time": 0.01,
                "status": "OK",
            }
        ]
        ecu_reports = {
            "BCM": {
                "ecu": "BCM",
                "activities": [
                    {
                        "timestamp": 0.123,
                        "service_name": "ReadDataByIdentifier",
                        "display_name": "Read VIN",
                        "request": "22 F1 90",
                        "response": "62 F1 90",
                        "response_time": 0.01,
                        "status": "OK",
                        "response_pending": 0,
                    }
                ],
            }
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            output_file = Path(tmpdir) / "log_analyzer.xlsx"

            self.assertTrue(
                export_workbook(
                    summary,
                    ecu_reports,
                    transactions,
                    output_file,
                )
            )

            workbook = load_workbook(output_file)

        self.assertEqual(workbook.sheetnames, ["Summary", "ALL_ECUS", "BCM"])
        self.assertEqual(workbook["Summary"]["A2"].value, "BCM")
        self.assertEqual(workbook["ALL_ECUS"]["A2"].value, "BCM")
        self.assertEqual(
            [cell.value for cell in workbook["ALL_ECUS"][1]],
            [
                "ECU",
                "Time",
                "Activity",
                "Request",
                "Response",
                "RT (ms)",
                "Status",
            ],
        )
        self.assertEqual(workbook["BCM"]["D2"].value, "Read VIN")

        for sheet_name in ["Summary", "ALL_ECUS", "BCM"]:
            sheet = workbook[sheet_name]
            header = sheet["A1"]
            self.assertTrue(header.font.bold)
            self.assertTrue(_cell_color(header).endswith(COLOR_HEADER))
            self.assertEqual(header.alignment.horizontal, "center")
            self.assertEqual(header.border.left.style, "thin")
            self.assertEqual(sheet["A2"].border.left.style, "thin")
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)


if __name__ == "__main__":
    unittest.main()
