import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from core.coding_value import (
    export_coding_value_rows_to_json,
    load_coding_value_rows,
    load_coding_value_rows_from_json,
)


class CodingValueExcelTests(unittest.TestCase):

    def test_loads_parameter_rows_and_decodes_payload_values(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CDS Coding"
        sheet.append([
            "DID Number (hex)",
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "0xF112",
            "Vehicle Name",
            3,
            0,
            8,
            "0x0=Unsupported\n0xFF=Invalid",
        ])
        sheet.append([
            None,
            "Steering Wheel Position",
            4,
            4,
            4,
            "0x1=Left hand drive (LHD)\n0x2=Right hand drive (RHD)",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            rows = load_coding_value_rows(
                path,
                "62 F1 12 FF 21",
            )

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].parameter, "Vehicle Name")
        self.assertEqual(rows[0].byte_pos, "3")
        self.assertEqual(rows[0].bit_pos, "0")
        self.assertEqual(rows[0].bit_length, "8")
        self.assertEqual(rows[0].raw_value, "0xFF")
        self.assertEqual(rows[0].decoded_value, "Invalid")
        self.assertEqual(
            [option.label for option in rows[0].decoded_options],
            ["Unsupported", "Invalid"],
        )

        self.assertEqual(rows[1].raw_value, "0x2")
        self.assertEqual(rows[1].decoded_value, "Right hand drive (RHD)")

    def test_loads_dollar_header_columns_as_parameters(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Field", "$Vehicle Name"])
        sheet.append(["Byte Pos", 3])
        sheet.append(["Bit Pos", 0])
        sheet.append(["Bit Length", 8])
        sheet.append(["Decoded Value", "0xFF=Invalid"])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            rows = load_coding_value_rows(
                path,
                "62 F1 12 FF",
            )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].parameter, "Vehicle Name")
        self.assertEqual(rows[0].raw_value, "0xFF")
        self.assertEqual(rows[0].decoded_value, "Invalid")

    def test_prefers_dollar_prefixed_parameter_table_columns(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "$Parameter",
            "$BytePos (from 0)",
            "$BitPos",
            "$BitLength",
            "$MethodType",
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            3,
            0,
            8,
            "0xFF=Invalid",
            "Wrong Parameter",
            0,
            0,
            8,
            "0x62=Wrong",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            rows = load_coding_value_rows(
                path,
                "62 F1 12 FF",
            )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].parameter, "Vehicle Name")
        self.assertEqual(rows[0].byte_pos, "3")
        self.assertEqual(rows[0].raw_value, "0xFF")
        self.assertEqual(rows[0].decoded_value, "Invalid")
    def test_method_type_multiline_values_become_individual_options(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "$MethodType",
        ])
        sheet.append([
            "Vehicle Variant",
            14,
            0,
            8,
            "\n0x0=Unsupported\n0x1= Comfort\n0x2= Premium\n"
            "0x3= Earth\n0x4= Wind\n0x5=  Sky\n"
            "0x6=  HerioGreen\n0x7=  Reserved\n"
            "0x8 = Reserved\n0x9 = Reserved\n"
            "0xA = Reserved\n0xF=Invalid",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            rows = load_coding_value_rows(
                path,
                "00 " * 14 + "02",
            )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].decoded_value, "Premium")
        self.assertEqual(
            [option.label for option in rows[0].decoded_options],
            [
                "Unsupported",
                "Comfort",
                "Premium",
                "Earth",
                "Wind",
                "Sky",
                "HerioGreen",
                "Reserved",
                "Reserved",
                "Reserved",
                "Reserved",
                "Invalid",
            ],
        )

    def test_method_type_accepts_colon_separator(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "$Parameter",
            "$BytePos (from 0)",
            "$BitPos",
            "$BitLength",
            "$MethodType",
        ])
        sheet.append([
            "Body Color",
            14,
            0,
            8,
            "0x0: Unsupported (NA)\n"
            "0x1: Brahminy White\n"
            "0x2: De Sat Silver\n"
            "0x3: Neptune Grey",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            rows = load_coding_value_rows(
                path,
                "00 " * 14 + "02",
            )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].decoded_value, "De Sat Silver")
        self.assertEqual(
            [option.label for option in rows[0].decoded_options],
            [
                "Unsupported (NA)",
                "Brahminy White",
                "De Sat Silver",
                "Neptune Grey",
            ],
        )

    def test_method_type_accepts_hyphen_separator(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "$Parameter",
            "$BytePos (from 0)",
            "$BitPos",
            "$BitLength",
            "$MethodType",
        ])
        sheet.append([
            "Body Color",
            14,
            0,
            8,
            "0x01 - Brahminy White\n"
            "0X02 - De Sat Silver\n"
            "0x03 - Neptune Grey",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            rows = load_coding_value_rows(
                path,
                "00 " * 14 + "02",
            )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].decoded_value, "De Sat Silver")
        self.assertEqual(
            [option.label for option in rows[0].decoded_options],
            [
                "Brahminy White",
                "De Sat Silver",
                "Neptune Grey",
            ],
        )
    def test_bit_length_text_is_parsed_for_bit_field_decode(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Nibble Value",
            39,
            0,
            "4.0",
            "0x8=Eight",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            rows = load_coding_value_rows(
                path,
                "00 " * 39 + "08",
            )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].raw_value, "0x8")
        self.assertEqual(rows[0].decoded_value, "Eight")

    def test_json_round_trip_preserves_rows_and_options(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Variant",
            2,
            0,
            8,
            "0x01=Eco\n0x02=Sport",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "coding.xlsx"
            json_dir = Path(tmpdir) / "json"
            workbook.save(excel_path)

            rows = load_coding_value_rows(excel_path, "00 00 02")
            json_path = export_coding_value_rows_to_json(
                excel_path,
                rows,
                json_dir,
            )
            loaded = load_coding_value_rows_from_json(json_path)

        self.assertEqual(json_path.name, "coding.json")
        self.assertEqual(len(loaded), 1)
        self.assertEqual(loaded[0].parameter, "Vehicle Variant")
        self.assertEqual(loaded[0].byte_pos, "2")
        self.assertEqual(loaded[0].raw_value, "0x02")
        self.assertEqual(loaded[0].decoded_value, "Sport")
        self.assertEqual(
            [(option.raw_value, option.label) for option in loaded[0].decoded_options],
            [("0x01", "Eco"), ("0x02", "Sport")],
        )

    def test_skips_blank_parameter_rows_and_keeps_unmapped_raw_value(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([None, 0, 0, 8, "0x00=Blank"])
        sheet.append(["Unmapped Byte", 1.0, 0.0, 8.0, "0x01=One"])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            rows = load_coding_value_rows(path, "00 34")

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].parameter, "Unmapped Byte")
        self.assertEqual(rows[0].byte_pos, "1")
        self.assertEqual(rows[0].bit_pos, "0")
        self.assertEqual(rows[0].bit_length, "8")
        self.assertEqual(rows[0].raw_value, "0x34")
        self.assertEqual(rows[0].decoded_value, "0x34")

    def test_extracts_multi_byte_raw_value_in_payload_order(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append(["Two Byte Field", 1, 0, 16, "0x1234=Known"])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            rows = load_coding_value_rows(path, "AA 12 34")

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].raw_value, "0x1234")
        self.assertEqual(rows[0].decoded_value, "Known")
if __name__ == "__main__":
    unittest.main()
