import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from core.coding_value import CodingDefinition, CodingValueOption, CodingValueRow
from core.coding_value_definition import (
    export_coding_definition_to_json,
    load_coding_definition_from_excel,
    load_coding_definition_from_json,
)


class CodingValueDefinitionTests(unittest.TestCase):

    def test_loads_excel_as_coding_definition(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append(["Vehicle Name", 1.0, 0.0, 8.0, "0x01=Eco\n0x02=Sport"])

        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "MHU_Coding.xlsx"
            workbook.save(excel_path)

            definition = load_coding_definition_from_excel(excel_path, "00 02")

        self.assertIsInstance(definition, CodingDefinition)
        self.assertEqual(definition.name, "MHU_Coding")
        self.assertEqual(definition.source_file, "MHU_Coding.xlsx")
        self.assertEqual(definition.rows[0].parameter, "Vehicle Name")
        self.assertEqual(definition.rows[0].byte_pos, "1")
        self.assertEqual(definition.rows[0].bit_pos, "0")
        self.assertEqual(definition.rows[0].bit_length, "8")
        self.assertEqual(definition.rows[0].raw_value, "0x02")
        self.assertEqual(definition.rows[0].decoded_value, "Sport")
        self.assertIsInstance(definition.rows, tuple)

    def test_json_round_trip_preserves_rows_options_and_definition_model(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append(["Vehicle Variant", 2, 0, 8, "0x01=Eco\n0x02=Sport"])

        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "coding.xlsx"
            json_dir = Path(tmpdir) / "json"
            workbook.save(excel_path)

            excel_definition = load_coding_definition_from_excel(
                excel_path,
                "00 00 02",
            )
            json_path = export_coding_definition_to_json(
                excel_definition,
                json_dir,
            )
            json_definition = load_coding_definition_from_json(json_path)

        self.assertEqual(json_path.name, "coding.json")
        self.assertEqual(json_definition.name, "coding")
        self.assertEqual(json_definition.rows, excel_definition.rows)
        self.assertEqual(
            [(option.raw_value, option.label) for option in json_definition.rows[0].decoded_options],
            [("0x01", "Eco"), ("0x02", "Sport")],
        )

    def test_loads_old_json_schema_as_definition(self):
        payload = {
            "source_file": "coding.xlsx",
            "source_path": "C:/tmp/coding.xlsx",
            "rows": [
                {
                    "parameter": "Vehicle Name",
                    "byte_pos": "13",
                    "bit_pos": "0",
                    "bit_length": "8",
                    "raw_value": "0x09",
                    "decoded_value": "VF7NP",
                    "decoded_options": [
                        {"raw_value": "0x03", "label": "VF3"},
                        {"raw_value": "0x09", "label": "VF7NP"},
                    ],
                },
            ],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            json_path = Path(tmpdir) / "legacy.json"
            json_path.write_text(json.dumps(payload), encoding="utf-8")

            definition = load_coding_definition_from_json(json_path)

        self.assertEqual(definition.name, "coding")
        self.assertEqual(definition.source_file, "coding.xlsx")
        self.assertEqual(definition.source_path, "C:/tmp/coding.xlsx")
        self.assertEqual(definition.rows[0].decoded_value, "VF7NP")
        self.assertEqual(
            [(option.raw_value, option.label) for option in definition.rows[0].decoded_options],
            [("0x03", "VF3"), ("0x09", "VF7NP")],
        )

    def test_json_without_metadata_uses_json_file_stem_for_name(self):
        payload = {
            "rows": [
                {
                    "parameter": "Method Type",
                    "byte_pos": "14",
                    "bit_pos": "0",
                    "bit_length": "8",
                    "raw_value": "",
                    "decoded_value": "",
                    "decoded_options": [],
                },
            ],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            json_path = Path(tmpdir) / "VCU.json"
            json_path.write_text(json.dumps(payload), encoding="utf-8")

            definition = load_coding_definition_from_json(json_path)

        self.assertEqual(definition.name, "VCU")
        self.assertEqual(definition.source_file, "")
        self.assertEqual(definition.rows[0].parameter, "Method Type")

    def test_unmapped_raw_value_fallback_is_preserved(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append(["Unmapped Byte", 1, 0, 8, "0x01=One"])

        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "coding.xlsx"
            workbook.save(excel_path)

            definition = load_coding_definition_from_excel(excel_path, "00 34")

        self.assertEqual(definition.rows[0].raw_value, "0x34")
        self.assertEqual(definition.rows[0].decoded_value, "0x34")

    def test_raw_format_compatibility_for_multi_byte_and_nibble(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append(["Two Byte Field", 1, 0, 16, "0x1234=Known"])
        sheet.append(["Nibble Field", 3, 0, "4.0", "0x8=Eight"])

        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "coding.xlsx"
            workbook.save(excel_path)

            definition = load_coding_definition_from_excel(
                excel_path,
                "AA 12 34 08",
            )

        self.assertEqual(definition.rows[0].raw_value, "0x1234")
        self.assertEqual(definition.rows[0].decoded_value, "Known")
        self.assertEqual(definition.rows[1].raw_value, "0x8")
        self.assertEqual(definition.rows[1].decoded_value, "Eight")

    def test_two_definitions_do_not_share_mutable_row_state(self):
        row_a = CodingValueRow(
            parameter="A",
            byte_pos="0",
            bit_pos="0",
            bit_length="8",
            raw_value="",
            decoded_value="",
            decoded_options=(CodingValueOption(raw_value="0x01", label="One"),),
        )
        row_b = CodingValueRow(
            parameter="B",
            byte_pos="1",
            bit_pos="0",
            bit_length="8",
            raw_value="",
            decoded_value="",
            decoded_options=(),
        )

        definition_a = CodingDefinition(name="A", rows=(row_a,))
        definition_b = CodingDefinition(name="B", rows=(row_b,))

        self.assertIsNot(definition_a.rows, definition_b.rows)
        self.assertEqual(definition_a.rows[0].parameter, "A")
        self.assertEqual(definition_b.rows[0].parameter, "B")
        with self.assertRaises(AttributeError):
            definition_a.name = "changed"


if __name__ == "__main__":
    unittest.main()
