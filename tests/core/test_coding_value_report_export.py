from datetime import datetime
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from core.coding_value_report_export import (
    CODING_VALUE_REPORT_HEADERS,
    CodingValueReportData,
    CodingValueReportRow,
    export_coding_value_report,
)
from core.excel_styles import COLOR_HEADER, COLOR_WARNING
from core.services.coding_value_report_service import CodingValueReportService


def _excel_color(cell):
    value = cell.fill.fgColor.rgb
    return "" if value is None else str(value).upper()


class CodingValueReportExportTests(unittest.TestCase):

    def test_exporter_creates_workbook_with_preserved_headers_and_values(self):
        report_data = CodingValueReportData(
            sheet_name="Coding value_62 F1 08",
            headers=CODING_VALUE_REPORT_HEADERS,
            rows=(
                CodingValueReportRow(
                    parameter="Vehicle Name",
                    byte_pos="13",
                    bit_pos="0",
                    bit_length="8",
                    raw_value="01",
                    decoded_value="VF5",
                    raw_value_before="09",
                    decoded_value_before="VF7NP",
                    result="No-M",
                ),
                CodingValueReportRow(
                    parameter="Unchanged Byte",
                    byte_pos="14",
                    bit_pos="0",
                    bit_length="8",
                    raw_value="03",
                    decoded_value="VF3",
                    raw_value_before="03",
                    decoded_value_before="VF3",
                    result="MATCH",
                ),
            ),
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            output_file = Path(tmpdir) / "coding_report.xlsx"

            exported_path = export_coding_value_report(report_data, output_file)

            self.assertEqual(exported_path, output_file)
            self.assertTrue(output_file.exists())
            workbook = load_workbook(output_file)
            sheet = workbook["Coding value_62 F1 08"]

        self.assertEqual(
            [cell.value for cell in sheet[1]],
            CODING_VALUE_REPORT_HEADERS,
        )
        self.assertEqual(sheet.cell(row=2, column=1).value, "Vehicle Name")
        self.assertEqual(sheet.cell(row=2, column=5).value, "01")
        self.assertEqual(sheet.cell(row=2, column=6).value, "VF5")
        self.assertEqual(sheet.cell(row=2, column=7).value, "09")
        self.assertEqual(sheet.cell(row=2, column=8).value, "VF7NP")
        self.assertEqual(sheet.cell(row=2, column=9).value, "No-M")
        self.assertEqual(sheet.cell(row=3, column=9).value, "MATCH")

    def test_exporter_preserves_report_styles(self):
        report_data = CodingValueReportData(
            sheet_name="Coding value_62 F1 08",
            headers=CODING_VALUE_REPORT_HEADERS,
            rows=(
                CodingValueReportRow(
                    parameter="Changed Byte",
                    byte_pos="13",
                    bit_pos="0",
                    bit_length="8",
                    raw_value="01",
                    decoded_value="VF5",
                    raw_value_before="09",
                    decoded_value_before="VF7NP",
                    result="No-M",
                ),
                CodingValueReportRow(
                    parameter="Unchanged Byte",
                    byte_pos="14",
                    bit_pos="0",
                    bit_length="8",
                    raw_value="03",
                    decoded_value="VF3",
                    raw_value_before="03",
                    decoded_value_before="VF3",
                    result="MATCH",
                ),
            ),
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            output_file = Path(tmpdir) / "coding_report.xlsx"
            export_coding_value_report(report_data, output_file)
            workbook = load_workbook(output_file)
            sheet = workbook["Coding value_62 F1 08"]

        self.assertTrue(sheet["A1"].font.bold)
        self.assertTrue(_excel_color(sheet["A1"]).endswith(COLOR_HEADER))
        self.assertEqual(sheet["A1"].alignment.horizontal, "center")
        self.assertEqual(sheet["A1"].border.left.style, "thin")
        self.assertEqual(sheet["A2"].border.left.style, "thin")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        for cell in sheet[2]:
            self.assertEqual(cell.fill.fill_type, "solid")
            self.assertTrue(cell.fill.fgColor.rgb.upper().endswith(COLOR_WARNING))
            self.assertEqual(cell.border.left.style, "thin")
        for cell in sheet[3]:
            self.assertNotEqual(cell.fill.fill_type, "solid")
            self.assertEqual(cell.border.left.style, "thin")

    def test_service_returns_default_path_and_delegates_export(self):
        calls = []

        def exporter(report_data, output_file):
            calls.append((report_data, Path(output_file)))
            Path(output_file).write_text("exported", encoding="utf-8")
            return Path(output_file)

        report_data = CodingValueReportData(
            sheet_name="Coding value_62 F1 08",
            headers=CODING_VALUE_REPORT_HEADERS,
            rows=(),
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            service = CodingValueReportService(
                report_root=Path(tmpdir) / "report",
                exporter=exporter,
                now_provider=lambda: datetime(2026, 10, 2, 3, 4, 5),
            )
            default_path = service.default_report_path()
            output_file = Path(tmpdir) / "custom.xlsx"

            exported_path = service.export_report(report_data, output_file)

        self.assertEqual(default_path.parent.name, "02-10-2026")
        self.assertEqual(
            default_path.name,
            "EEIV_report_Coding value_20261002_030405.xlsx",
        )
        self.assertEqual(exported_path, output_file)
        self.assertEqual(calls, [(report_data, output_file)])


if __name__ == "__main__":
    unittest.main()
