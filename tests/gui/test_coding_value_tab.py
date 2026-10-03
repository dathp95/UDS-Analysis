import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QEvent, Qt
from PySide6.QtGui import QFocusEvent, QKeySequence
from PySide6.QtWidgets import QApplication, QHBoxLayout, QHeaderView, QPlainTextEdit, QStackedWidget, QVBoxLayout

from config.paths import CONFIG_DIR, EXPORT_CODING_FILES_DIR
from core.coding_value import CodingDefinition, CodingValueOption, CodingValueRow
from core.coding_value_workspace import CodingEcuWorkspace
from core.coding_value_workspace_store import (
    CodingWorkspaceState,
    CodingWorkspaceStore,
)
from core.coding_value_workspace_snapshot import CodingWorkspaceSnapshot
from core.excel_styles import COLOR_HEADER, COLOR_WARNING
from core.crc import calculate_crc8_sae_j1850
from core.coding_value_definition import export_coding_definition_to_json
from core.services.coding_value_report_service import CodingValueReportService
from gui.themes.theme_manager import ThemeManager
from gui.tabs.coding_value_tab import CodingValueTab
from gui.widgets.coding_value.ecu_workspace_sidebar import EcuWorkspaceSidebar
from gui.widgets.coding_value.coding_value_panel import CodingValuePanel
from gui.widgets.coding_value.coding_value_table import CodingValueTable


def _excel_color(cell):
    value = cell.fill.fgColor.rgb
    return "" if value is None else str(value).upper()


class CodingValueTabTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.application = QApplication.instance() or QApplication([])

    def setUp(self):
        self.export_dir_handle = tempfile.TemporaryDirectory()
        self.addCleanup(self.export_dir_handle.cleanup)
        self.export_dir_patch = patch(
            "gui.widgets.coding_value.coding_value_panel.EXPORT_CODING_FILES_DIR",
            Path(self.export_dir_handle.name) / "config" / "export_coding_files",
        )
        self.export_dir_patch.start()
        self.addCleanup(self.export_dir_patch.stop)
        self.workspace_store_handle = tempfile.TemporaryDirectory()
        self.addCleanup(self.workspace_store_handle.cleanup)

    def create_workspace_store(self):
        return CodingWorkspaceStore(
            Path(self.workspace_store_handle.name) / "coding_value_workspaces.json"
        )

    def create_tab(self, workspace_store=None):
        return CodingValueTab(
            workspace_store=workspace_store or self.create_workspace_store()
        )

    def create_definition_file(self, directory: Path, name: str, parameter: str) -> Path:
        directory.mkdir(parents=True, exist_ok=True)
        definition = CodingDefinition(
            name=name,
            rows=(
                CodingValueRow(
                    parameter=parameter,
                    byte_pos="0",
                    bit_pos="0",
                    bit_length="8",
                    raw_value="",
                    decoded_value="",
                    decoded_options=(),
                ),
            ),
            source_file=f"{name}.xlsx",
            source_path=f"config/Coding/{name}.xlsx",
        )
        return export_coding_definition_to_json(definition, directory / f"{name}.json")

    def assert_same_path(self, left, right):
        self.assertEqual(
            Path(str(left)).resolve(strict=False),
            Path(str(right)).resolve(strict=False),
        )

    def active_sidebar_item_ids(self, tab):
        return [
            workspace_id
            for workspace_id, item in tab.workspace_sidebar._item_by_id.items()
            if item._active
        ]

    def load_payload_on_workspace(
            self,
            tab,
            workspace_id,
            definition_name,
            parameter,
            payload,
        ):
        tab.workspace_sidebar.select_workspace(workspace_id)
        panel = tab._panel_by_workspace_id[workspace_id]
        definition_path = self.create_definition_file(
            Path(self.export_dir_handle.name) / "config" / "export_coding_files",
            definition_name,
            parameter,
        )
        self.assertTrue(panel.load_coding_definition_file(str(definition_path)))
        panel.txt_coding_value.setPlainText(payload)
        panel.encode_coding_payload()
        return panel

    def test_coding_value_tab_displays_sidebar_and_initial_stack_panel(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)

        self.assertIsInstance(tab.workspace_sidebar, EcuWorkspaceSidebar)
        self.assertIsInstance(tab.coding_value_panel, CodingValuePanel)
        self.assertIsInstance(tab.workspace_stack, QStackedWidget)
        self.assertIsInstance(tab.workspace_row.layout(), QHBoxLayout)
        self.assertEqual(tab.workspace_row.layout().indexOf(tab.workspace_sidebar), 0)
        self.assertEqual(tab.workspace_row.layout().indexOf(tab.workspace_stack), 1)
        self.assertEqual(tab.workspace_stack.count(), 1)
        self.assertEqual(len(tab.findChildren(CodingValuePanel)), 1)
        self.assertIs(tab.workspace_stack.currentWidget(), tab.coding_value_panel)

    def test_coding_value_tab_creates_default_workspace_and_panel_mapping(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)

        self.assertEqual(tab.workspace_sidebar.workspace_count(), 1)
        workspace = tab.workspace_sidebar.workspaces()[0]
        self.assertEqual(workspace.name, "ECU 1")
        self.assertEqual(tab.workspace_sidebar.active_workspace_id(), workspace.id)
        self.assertEqual(set(tab._panel_by_workspace_id), {workspace.id})
        self.assertIs(tab.active_coding_value_panel(), tab._panel_by_workspace_id[workspace.id])

    def test_coding_value_tab_restores_saved_workspaces_without_extra_default(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            store = CodingWorkspaceStore(Path(tmpdir) / "workspaces.json")
            mhu = CodingEcuWorkspace("mhu-id", "MHU")
            vcu = CodingEcuWorkspace("vcu-id", "VCU")
            bcm = CodingEcuWorkspace("bcm-id", "BCM")
            store.save(CodingWorkspaceState((mhu, vcu, bcm), "vcu-id"))

            tab = self.create_tab(store)
            self.addCleanup(tab.deleteLater)

        self.assertEqual(
            [workspace.name for workspace in tab.workspace_sidebar.workspaces()],
            ["MHU", "VCU", "BCM"],
        )
        self.assertEqual(tab.workspace_stack.count(), 3)
        self.assertEqual(set(tab._panel_by_workspace_id), {"mhu-id", "vcu-id", "bcm-id"})
        self.assertEqual(tab.workspace_sidebar.active_workspace_id(), "vcu-id")
        self.assertIs(tab.workspace_stack.currentWidget(), tab._panel_by_workspace_id["vcu-id"])

    def test_coding_value_tab_restores_valid_empty_workspace_config(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            store = CodingWorkspaceStore(Path(tmpdir) / "workspaces.json")
            store.save(CodingWorkspaceState((), None))

            tab = self.create_tab(store)
            self.addCleanup(tab.deleteLater)

        self.assertEqual(tab.workspace_sidebar.workspace_count(), 0)
        self.assertEqual(tab.workspace_stack.count(), 0)
        self.assertIsNone(tab.active_coding_value_panel())

    def test_coding_value_tab_uses_default_for_corrupt_workspace_config(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            config_file = Path(tmpdir) / "workspaces.json"
            config_file.write_text("{ bad json", encoding="utf-8")

            tab = self.create_tab(CodingWorkspaceStore(config_file))
            self.addCleanup(tab.deleteLater)

            self.assertEqual(tab.workspace_sidebar.workspace_count(), 1)
            self.assertEqual(tab.workspace_sidebar.workspaces()[0].name, "ECU 1")
            self.assertEqual(config_file.read_text(encoding="utf-8"), "{ bad json")

    def test_workspace_add_rename_delete_and_selection_persist(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            store = CodingWorkspaceStore(Path(tmpdir) / "workspaces.json")
            tab = self.create_tab(store)
            self.addCleanup(tab.deleteLater)

            first_workspace = tab.workspace_sidebar.workspaces()[0]
            tab.workspace_sidebar.rename_workspace(first_workspace.id, "MHU")
            with patch(
                "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
                return_value=("VCU", True),
            ):
                tab.workspace_sidebar.btn_add.click()
            vcu_workspace = tab.workspace_sidebar.active_workspace()
            tab.workspace_sidebar.rename_workspace(vcu_workspace.id, "Vehicle Control")
            tab.workspace_sidebar.select_workspace(first_workspace.id)
            tab.workspace_sidebar.remove_workspace(vcu_workspace.id)

            reloaded = self.create_tab(store)
            self.addCleanup(reloaded.deleteLater)

        self.assertEqual(
            [(workspace.id, workspace.name) for workspace in reloaded.workspace_sidebar.workspaces()],
            [(first_workspace.id, "MHU")],
        )
        self.assertEqual(reloaded.workspace_sidebar.active_workspace_id(), first_workspace.id)

    def test_active_selection_persists_across_restart(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            store = CodingWorkspaceStore(Path(tmpdir) / "workspaces.json")
            mhu = CodingEcuWorkspace("mhu-id", "MHU")
            vcu = CodingEcuWorkspace("vcu-id", "VCU")
            store.save(CodingWorkspaceState((mhu, vcu), "mhu-id"))
            tab = self.create_tab(store)
            self.addCleanup(tab.deleteLater)

            tab.workspace_sidebar.select_workspace("vcu-id")
            reloaded = self.create_tab(store)
            self.addCleanup(reloaded.deleteLater)

        self.assertEqual(reloaded.workspace_sidebar.active_workspace_id(), "vcu-id")

    def test_coding_file_association_and_definition_restore_per_workspace(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmpdir_path = Path(tmpdir)
            mhu_json = self.create_definition_file(tmpdir_path, "mhu", "MHU Byte")
            vcu_json = self.create_definition_file(tmpdir_path, "vcu", "VCU Byte")
            store = CodingWorkspaceStore(tmpdir_path / "workspaces.json")
            store.save(CodingWorkspaceState((
                CodingEcuWorkspace("mhu-id", "MHU", str(mhu_json)),
                CodingEcuWorkspace("vcu-id", "VCU", str(vcu_json)),
            ), "vcu-id"))

            tab = self.create_tab(store)
            self.addCleanup(tab.deleteLater)

        mhu_panel = tab._panel_by_workspace_id["mhu-id"]
        vcu_panel = tab._panel_by_workspace_id["vcu-id"]
        self.assertEqual(mhu_panel._coding_definition.rows[0].parameter, "MHU Byte")
        self.assertEqual(vcu_panel._coding_definition.rows[0].parameter, "VCU Byte")
        self.assertEqual(mhu_panel.txt_coding_value.toPlainText(), "")
        self.assertEqual(vcu_panel.txt_coding_value.toPlainText(), "")

    def test_relative_coding_files_restore_distinct_panels_and_combo_selection(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            project_root = Path(tmpdir)
            export_dir = project_root / "config" / "export_coding_files"
            mhu_json = self.create_definition_file(export_dir, "mhu", "MHU Byte")
            vcu_json = self.create_definition_file(export_dir, "vcu", "VCU Byte")
            bcm_json = self.create_definition_file(export_dir, "bcm", "BCM Byte")
            store = CodingWorkspaceStore(
                project_root / "config" / "coding_value_workspaces.json",
                project_root=project_root,
            )
            store.save(CodingWorkspaceState((
                CodingEcuWorkspace(
                    "mhu-id",
                    "MHU",
                    "config/export_coding_files/mhu.json",
                ),
                CodingEcuWorkspace(
                    "vcu-id",
                    "VCU",
                    "config/export_coding_files/vcu.json",
                ),
                CodingEcuWorkspace(
                    "bcm-id",
                    "BCM",
                    "config/export_coding_files/bcm.json",
                ),
            ), "vcu-id"))

            with patch(
                "gui.widgets.coding_value.coding_value_panel.EXPORT_CODING_FILES_DIR",
                export_dir,
            ):
                tab = self.create_tab(store)
                self.addCleanup(tab.deleteLater)

        mhu_panel = tab._panel_by_workspace_id["mhu-id"]
        vcu_panel = tab._panel_by_workspace_id["vcu-id"]
        bcm_panel = tab._panel_by_workspace_id["bcm-id"]
        self.assertEqual(mhu_panel._coding_definition.rows[0].parameter, "MHU Byte")
        self.assertEqual(vcu_panel._coding_definition.rows[0].parameter, "VCU Byte")
        self.assertEqual(bcm_panel._coding_definition.rows[0].parameter, "BCM Byte")
        self.assert_same_path(mhu_panel.cmb_coding_json.currentData(), mhu_json)
        self.assert_same_path(vcu_panel.cmb_coding_json.currentData(), vcu_json)
        self.assert_same_path(bcm_panel.cmb_coding_json.currentData(), bcm_json)
        self.assertEqual(tab.workspace_sidebar.active_workspace_id(), "vcu-id")
        self.assertEqual(self.active_sidebar_item_ids(tab), ["vcu-id"])
        self.assertIs(tab.workspace_stack.currentWidget(), vcu_panel)
        self.assertEqual(
            [workspace.coding_file for workspace in tab.workspace_sidebar.workspaces()],
            [
                "config/export_coding_files/mhu.json",
                "config/export_coding_files/vcu.json",
                "config/export_coding_files/bcm.json",
            ],
        )

    def test_new_workspace_starts_without_copying_active_definition(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)
        mhu_workspace = tab.workspace_sidebar.workspaces()[0]
        definition_path = self.create_definition_file(
            Path(self.export_dir_handle.name) / "config" / "export_coding_files",
            "mhu",
            "MHU Byte",
        )
        mhu_panel = tab._panel_by_workspace_id[mhu_workspace.id]

        self.assertTrue(mhu_panel.load_coding_definition_file(str(definition_path)))
        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("VCU", True),
        ):
            tab.workspace_sidebar.btn_add.click()

        vcu_workspace = tab.workspace_sidebar.active_workspace()
        vcu_panel = tab._panel_by_workspace_id[vcu_workspace.id]
        self.assertEqual(mhu_workspace.coding_file, str(definition_path))
        self.assertEqual(vcu_workspace.coding_file, "")
        self.assertIsNotNone(mhu_panel._coding_definition)
        self.assertIsNone(vcu_panel._coding_definition)
        self.assert_same_path(mhu_panel.cmb_coding_json.currentData(), definition_path)
        self.assertEqual(vcu_panel.cmb_coding_json.currentIndex(), -1)

    def test_manual_save_restores_independent_workspace_snapshots_after_restart(self):
        store = self.create_workspace_store()
        tab = self.create_tab(store)
        self.addCleanup(tab.deleteLater)
        mhu_workspace = tab.workspace_sidebar.workspaces()[0]
        tab.workspace_sidebar.rename_workspace(mhu_workspace.id, "MHU")
        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("VCU", True),
        ):
            tab.workspace_sidebar.btn_add.click()
        vcu_workspace = tab.workspace_sidebar.active_workspace()

        mhu_panel = self.load_payload_on_workspace(
            tab, mhu_workspace.id, "mhu", "MHU Byte", "AA 11"
        )
        tab._save_workspace_snapshot(mhu_workspace.id)
        vcu_panel = self.load_payload_on_workspace(
            tab, vcu_workspace.id, "vcu", "VCU Byte", "BB 22"
        )
        tab._save_workspace_snapshot(vcu_workspace.id)
        reloaded = self.create_tab(store)
        self.addCleanup(reloaded.deleteLater)

        restored_mhu = reloaded._panel_by_workspace_id[mhu_workspace.id]
        restored_vcu = reloaded._panel_by_workspace_id[vcu_workspace.id]
        self.assertEqual(mhu_panel.txt_coding_value.toPlainText(), "AA 11")
        self.assertEqual(vcu_panel.txt_coding_value.toPlainText(), "BB 22")
        self.assertEqual(restored_mhu.txt_coding_value.toPlainText(), "AA 11")
        self.assertEqual(restored_mhu.txt_coding_preview.toPlainText(), "AA 11")
        self.assertEqual(restored_mhu.table.item(0, 0).text(), "MHU Byte")
        self.assertEqual(restored_vcu.txt_coding_value.toPlainText(), "BB 22")
        self.assertEqual(restored_vcu.txt_coding_preview.toPlainText(), "BB 22")
        self.assertEqual(restored_vcu.table.item(0, 0).text(), "VCU Byte")

    def test_unsaved_edit_does_not_survive_restart_but_stays_live_in_session(self):
        store = self.create_workspace_store()
        tab = self.create_tab(store)
        self.addCleanup(tab.deleteLater)
        mhu_workspace = tab.workspace_sidebar.workspaces()[0]
        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("VCU", True),
        ):
            tab.workspace_sidebar.btn_add.click()
        vcu_workspace = tab.workspace_sidebar.active_workspace()

        mhu_panel = self.load_payload_on_workspace(
            tab, mhu_workspace.id, "mhu", "MHU Byte", "AA"
        )
        tab._save_workspace_snapshot(mhu_workspace.id)
        mhu_panel.txt_coding_value.setPlainText("AB")
        mhu_panel.refresh_coding_preview()
        tab.workspace_sidebar.select_workspace(vcu_workspace.id)
        tab.workspace_sidebar.select_workspace(mhu_workspace.id)

        self.assertEqual(mhu_panel.txt_coding_preview.toPlainText(), "AB")

        reloaded = self.create_tab(store)
        self.addCleanup(reloaded.deleteLater)
        restored_mhu = reloaded._panel_by_workspace_id[mhu_workspace.id]

        self.assertEqual(restored_mhu.txt_coding_value.toPlainText(), "AA")
        self.assertEqual(restored_mhu.txt_coding_preview.toPlainText(), "AA")

    def test_manual_save_updated_state_survives_restart(self):
        store = self.create_workspace_store()
        tab = self.create_tab(store)
        self.addCleanup(tab.deleteLater)
        workspace = tab.workspace_sidebar.workspaces()[0]
        panel = self.load_payload_on_workspace(
            tab, workspace.id, "mhu", "MHU Byte", "AA"
        )
        tab._save_workspace_snapshot(workspace.id)

        panel.txt_coding_value.setPlainText("AB")
        panel.refresh_coding_preview()
        tab._save_workspace_snapshot(workspace.id)
        reloaded = self.create_tab(store)
        self.addCleanup(reloaded.deleteLater)

        restored_panel = reloaded._panel_by_workspace_id[workspace.id]
        self.assertEqual(restored_panel.txt_coding_value.toPlainText(), "AB")
        self.assertEqual(restored_panel.txt_coding_preview.toPlainText(), "AB")

    def test_snapshot_restore_preserves_baseline_and_reconstructs_check_state(self):
        store = self.create_workspace_store()
        tab = self.create_tab(store)
        self.addCleanup(tab.deleteLater)
        workspace = tab.workspace_sidebar.workspaces()[0]
        panel = self.load_payload_on_workspace(
            tab, workspace.id, "mhu", "MHU Byte", "37"
        )

        panel.table.fn_set_raw_value_at_row(0, "FF")
        panel.check_coding_value()
        panel.filter_no_match_rows()
        tab._save_workspace_snapshot(workspace.id)
        reloaded = self.create_tab(store)
        self.addCleanup(reloaded.deleteLater)

        restored_panel = reloaded._panel_by_workspace_id[workspace.id]
        self.assertEqual(restored_panel._payload_baseline_bytes, [0x37])
        self.assertEqual(restored_panel._payload_preview_bytes, [0xFF])
        self.assertEqual(restored_panel.table.item(0, 6).text(), "37")
        self.assertEqual(restored_panel.table.item(0, 8).text(), "No-M")
        self.assertEqual(restored_panel.txt_parameter_filter.text(), "No-M")
        self.assertIn("Total No-M: 1", restored_panel.txt_working_log.toPlainText())

    def test_rename_keeps_snapshot_by_workspace_id(self):
        store = self.create_workspace_store()
        tab = self.create_tab(store)
        self.addCleanup(tab.deleteLater)
        workspace = tab.workspace_sidebar.workspaces()[0]
        self.load_payload_on_workspace(
            tab, workspace.id, "mhu", "MHU Byte", "AA"
        )
        tab._save_workspace_snapshot(workspace.id)

        tab.workspace_sidebar.rename_workspace(workspace.id, "Main Head Unit")
        reloaded = self.create_tab(store)
        self.addCleanup(reloaded.deleteLater)

        restored_panel = reloaded._panel_by_workspace_id[workspace.id]
        self.assertEqual(
            reloaded.workspace_sidebar.workspace_by_id(workspace.id).name,
            "Main Head Unit",
        )
        self.assertEqual(restored_panel.txt_coding_preview.toPlainText(), "AA")

    def test_delete_workspace_removes_saved_snapshot(self):
        store = self.create_workspace_store()
        tab = self.create_tab(store)
        self.addCleanup(tab.deleteLater)
        first_workspace = tab.workspace_sidebar.workspaces()[0]
        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("VCU", True),
        ):
            tab.workspace_sidebar.btn_add.click()
        vcu_workspace = tab.workspace_sidebar.active_workspace()
        self.load_payload_on_workspace(
            tab, vcu_workspace.id, "vcu", "VCU Byte", "BB"
        )
        tab._save_workspace_snapshot(vcu_workspace.id)
        self.assertTrue(store.snapshot_path(vcu_workspace.id).exists())

        tab.workspace_sidebar.select_workspace(first_workspace.id)
        tab.workspace_sidebar.remove_workspace(vcu_workspace.id)

        self.assertFalse(store.snapshot_path(vcu_workspace.id).exists())
        reloaded = self.create_tab(store)
        self.addCleanup(reloaded.deleteLater)
        self.assertIsNone(reloaded.workspace_sidebar.workspace_by_id(vcu_workspace.id))

    def test_corrupt_snapshot_does_not_block_other_workspace_restore(self):
        store = self.create_workspace_store()
        mhu = CodingEcuWorkspace("mhu-id", "MHU")
        vcu = CodingEcuWorkspace("vcu-id", "VCU")
        store.save(CodingWorkspaceState((mhu, vcu), "vcu-id"))
        vcu_definition = self.create_definition_file(
            Path(self.export_dir_handle.name) / "config" / "export_coding_files",
            "vcu",
            "VCU Byte",
        )
        store.snapshot_path("mhu-id").parent.mkdir(parents=True, exist_ok=True)
        store.snapshot_path("mhu-id").write_text("{ bad json", encoding="utf-8")
        store.save_snapshot(
            "vcu-id",
            CodingWorkspaceSnapshot(
                coding_file=str(vcu_definition),
                payload_input="BB",
                payload_preview="BB",
                baseline_payload="BB",
                raw_values=(),
            ),
        )

        tab = self.create_tab(store)
        self.addCleanup(tab.deleteLater)

        self.assertIsNone(tab._panel_by_workspace_id["mhu-id"]._coding_definition)
        self.assertEqual(
            tab._panel_by_workspace_id["vcu-id"].txt_coding_preview.toPlainText(),
            "BB",
        )

    def test_save_non_active_workspace_does_not_switch_active_or_save_other_panel(self):
        store = self.create_workspace_store()
        tab = self.create_tab(store)
        self.addCleanup(tab.deleteLater)
        mhu_workspace = tab.workspace_sidebar.workspaces()[0]
        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("VCU", True),
        ):
            tab.workspace_sidebar.btn_add.click()
        vcu_workspace = tab.workspace_sidebar.active_workspace()
        mhu_panel = self.load_payload_on_workspace(
            tab, mhu_workspace.id, "mhu", "MHU Byte", "AA"
        )
        vcu_panel = self.load_payload_on_workspace(
            tab, vcu_workspace.id, "vcu", "VCU Byte", "BB"
        )
        tab._save_workspace_snapshot(vcu_workspace.id)
        tab.workspace_sidebar.select_workspace(vcu_workspace.id)
        mhu_panel.txt_coding_value.setPlainText("CC")
        mhu_panel.refresh_coding_preview()

        tab.workspace_sidebar.workspace_save_requested.emit(mhu_workspace.id)

        self.assertEqual(tab.workspace_sidebar.active_workspace_id(), vcu_workspace.id)
        self.assertEqual(vcu_panel.txt_coding_preview.toPlainText(), "BB")
        self.assertEqual(store.load_snapshot(mhu_workspace.id).payload_preview, "CC")
        self.assertEqual(store.load_snapshot(vcu_workspace.id).payload_preview, "BB")

    def test_dirty_marker_is_runtime_and_clears_after_manual_save(self):
        store = self.create_workspace_store()
        tab = self.create_tab(store)
        self.addCleanup(tab.deleteLater)
        workspace = tab.workspace_sidebar.workspaces()[0]
        item = tab.workspace_sidebar._item_by_id[workspace.id]

        panel = self.load_payload_on_workspace(
            tab, workspace.id, "mhu", "MHU Byte", "AA"
        )
        self.assertEqual(item.text(), "ECU 1 *")

        tab._save_workspace_snapshot(workspace.id)

        self.assertEqual(item.text(), "ECU 1")
        panel.txt_coding_value.setPlainText("AB")

        self.assertEqual(item.text(), "ECU 1 *")

    def test_theme_refresh_preserves_active_sidebar_visual(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)
        mhu_workspace = tab.workspace_sidebar.workspaces()[0]
        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("VCU", True),
        ):
            tab.workspace_sidebar.btn_add.click()
        vcu_workspace = tab.workspace_sidebar.active_workspace()

        tab.workspace_sidebar.select_workspace(mhu_workspace.id)
        tab.workspace_sidebar.select_workspace(vcu_workspace.id)
        tab.fn_refresh_theme()

        self.assertEqual(tab.workspace_sidebar.active_workspace_id(), vcu_workspace.id)
        self.assertEqual(self.active_sidebar_item_ids(tab), [vcu_workspace.id])

    def test_rename_preserves_active_definition_and_combo_selection(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)
        workspace = tab.workspace_sidebar.workspaces()[0]
        definition_path = self.create_definition_file(
            Path(self.export_dir_handle.name) / "config" / "export_coding_files",
            "mhu",
            "MHU Byte",
        )
        panel = tab._panel_by_workspace_id[workspace.id]
        self.assertTrue(panel.load_coding_definition_file(str(definition_path)))

        self.assertTrue(tab.workspace_sidebar.rename_workspace(workspace.id, "MHU"))

        self.assertEqual(tab.workspace_sidebar.active_workspace_id(), workspace.id)
        self.assertIs(tab.workspace_stack.currentWidget(), panel)
        self.assertEqual(workspace.coding_file, str(definition_path))
        self.assertEqual(panel._coding_definition.rows[0].parameter, "MHU Byte")
        self.assert_same_path(panel.cmb_coding_json.currentData(), definition_path)

    def test_delete_active_workspace_falls_back_and_persists_active(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            store = CodingWorkspaceStore(Path(tmpdir) / "workspaces.json")
            mhu = CodingEcuWorkspace("mhu-id", "MHU")
            vcu = CodingEcuWorkspace("vcu-id", "VCU")
            bcm = CodingEcuWorkspace("bcm-id", "BCM")
            store.save(CodingWorkspaceState((mhu, vcu, bcm), "vcu-id"))
            tab = self.create_tab(store)
            self.addCleanup(tab.deleteLater)

            self.assertTrue(tab.workspace_sidebar.remove_workspace("vcu-id"))
            saved_state = store.load()

        self.assertEqual(tab.workspace_sidebar.active_workspace_id(), "bcm-id")
        self.assertEqual(self.active_sidebar_item_ids(tab), ["bcm-id"])
        self.assertIs(
            tab.workspace_stack.currentWidget(),
            tab._panel_by_workspace_id["bcm-id"],
        )
        self.assertEqual(saved_state.active_workspace_id, "bcm-id")

    def test_delete_non_active_workspace_keeps_active_selection(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            store = CodingWorkspaceStore(Path(tmpdir) / "workspaces.json")
            mhu = CodingEcuWorkspace("mhu-id", "MHU")
            vcu = CodingEcuWorkspace("vcu-id", "VCU")
            bcm = CodingEcuWorkspace("bcm-id", "BCM")
            store.save(CodingWorkspaceState((mhu, vcu, bcm), "mhu-id"))
            tab = self.create_tab(store)
            self.addCleanup(tab.deleteLater)

            self.assertTrue(tab.workspace_sidebar.remove_workspace("bcm-id"))
            saved_state = store.load()

        self.assertEqual(tab.workspace_sidebar.active_workspace_id(), "mhu-id")
        self.assertEqual(self.active_sidebar_item_ids(tab), ["mhu-id"])
        self.assertIs(
            tab.workspace_stack.currentWidget(),
            tab._panel_by_workspace_id["mhu-id"],
        )
        self.assertEqual(saved_state.active_workspace_id, "mhu-id")

    def test_missing_coding_file_keeps_workspace_and_recorded_path(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            missing_path = Path(tmpdir) / "missing.json"
            store = CodingWorkspaceStore(Path(tmpdir) / "workspaces.json")
            store.save(CodingWorkspaceState((
                CodingEcuWorkspace("mhu-id", "MHU", str(missing_path)),
            ), "mhu-id"))

            tab = self.create_tab(store)
            self.addCleanup(tab.deleteLater)

        workspace = tab.workspace_sidebar.workspaces()[0]
        self.assertEqual(workspace.coding_file, str(missing_path))
        self.assertIsNone(tab._panel_by_workspace_id["mhu-id"]._coding_definition)

    def test_panel_coding_file_change_is_saved_with_captured_workspace_id(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmpdir_path = Path(tmpdir)
            definition_path = self.create_definition_file(tmpdir_path, "mhu", "MHU Byte")
            store = CodingWorkspaceStore(tmpdir_path / "workspaces.json")
            tab = self.create_tab(store)
            self.addCleanup(tab.deleteLater)
            workspace = tab.workspace_sidebar.workspaces()[0]
            panel = tab._panel_by_workspace_id[workspace.id]

            self.assertTrue(panel.load_coding_definition_file(str(definition_path)))
            reloaded = self.create_tab(store)
            self.addCleanup(reloaded.deleteLater)

        reloaded_workspace = reloaded.workspace_sidebar.workspaces()[0]
        self.assertEqual(reloaded_workspace.coding_file, str(definition_path))
        self.assertEqual(
            reloaded._panel_by_workspace_id[reloaded_workspace.id]._coding_definition.rows[0].parameter,
            "MHU Byte",
        )

    def test_imported_excel_persists_exported_json_definition_path(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            0,
            0,
            8,
            "0x01=One",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            tmpdir_path = Path(tmpdir)
            excel_path = tmpdir_path / "mhu.xlsx"
            workbook.save(excel_path)
            store = CodingWorkspaceStore(tmpdir_path / "workspaces.json")
            tab = self.create_tab(store)
            self.addCleanup(tab.deleteLater)
            workspace = tab.workspace_sidebar.workspaces()[0]
            tab.coding_value_panel.file_path.setText(str(excel_path))

            tab.coding_value_panel.import_coding_value()

            reloaded = self.create_tab(store)
            self.addCleanup(reloaded.deleteLater)

        reloaded_workspace = reloaded.workspace_sidebar.workspaces()[0]
        self.assertEqual(reloaded_workspace.id, workspace.id)
        self.assertTrue(reloaded_workspace.coding_file.endswith("mhu.json"))
        self.assertNotEqual(reloaded_workspace.coding_file, str(excel_path))
        self.assertEqual(
            reloaded.coding_value_panel._coding_definition.rows[0].parameter,
            "Vehicle Name",
        )

    def test_transient_payload_check_filter_and_log_are_not_persisted(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmpdir_path = Path(tmpdir)
            definition_path = self.create_definition_file(tmpdir_path, "mhu", "MHU Byte")
            store = CodingWorkspaceStore(tmpdir_path / "workspaces.json")
            tab = self.create_tab(store)
            self.addCleanup(tab.deleteLater)
            panel = tab.coding_value_panel
            panel.load_coding_definition_file(str(definition_path))
            panel.txt_coding_value.setPlainText("AA")
            panel.encode_coding_payload()
            panel.table.fn_set_raw_value_at_row(0, "BB")
            panel.check_coding_value()
            panel.filter_no_match_rows()

            reloaded = self.create_tab(store)
            self.addCleanup(reloaded.deleteLater)

        reloaded_panel = reloaded.coding_value_panel
        self.assertEqual(reloaded_panel._coding_definition.rows[0].parameter, "MHU Byte")
        self.assertEqual(reloaded_panel.txt_coding_value.toPlainText(), "")
        self.assertEqual(reloaded_panel.txt_coding_preview.toPlainText(), "")
        self.assertEqual(reloaded_panel.txt_parameter_filter.text(), "")
        self.assertEqual(reloaded_panel.txt_working_log.toPlainText(), "")
        self.assertFalse(reloaded_panel._has_check_results())

    def test_add_button_creates_independent_panel_before_selecting_workspace(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)
        events = []
        tab.workspace_sidebar.workspace_added.connect(
            lambda workspace: events.append(("added", workspace.id, tab._panel_by_workspace_id.get(workspace.id)))
        )
        tab.workspace_sidebar.workspace_selected.connect(
            lambda workspace_id: events.append(("selected", workspace_id, tab._panel_by_workspace_id.get(workspace_id)))
        )

        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("ECU 2", True),
        ):
            tab.workspace_sidebar.btn_add.click()

        self.assertEqual(tab.workspace_sidebar.workspace_count(), 2)
        self.assertEqual(tab.workspace_stack.count(), 2)
        self.assertEqual(len(tab.findChildren(CodingValuePanel)), 2)
        active_workspace = tab.workspace_sidebar.active_workspace()
        self.assertIs(tab.workspace_stack.currentWidget(), tab._panel_by_workspace_id[active_workspace.id])
        self.assertEqual([event[0] for event in events], ["added", "selected"])
        self.assertIsNotNone(events[0][2])
        self.assertIs(events[1][2], events[0][2])

    def test_rename_workspace_does_not_recreate_panel(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)
        workspace = tab.workspace_sidebar.workspaces()[0]
        original_panel = tab._panel_by_workspace_id[workspace.id]

        self.assertTrue(tab.workspace_sidebar.rename_workspace(workspace.id, "Main Head Unit"))

        self.assertIs(tab._panel_by_workspace_id[workspace.id], original_panel)
        self.assertIs(tab.coding_value_panel, original_panel)

    def test_coding_value_tab_crc_routing_uses_active_panel(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)
        first_workspace = tab.workspace_sidebar.workspaces()[0]
        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("ECU 2", True),
        ):
            tab.workspace_sidebar.btn_add.click()
        second_workspace = tab.workspace_sidebar.active_workspace()
        first_panel = tab._panel_by_workspace_id[first_workspace.id]
        second_panel = tab._panel_by_workspace_id[second_workspace.id]

        with patch.object(first_panel, "fn_set_crc_value", return_value=True) as first_set_crc:
            with patch.object(second_panel, "fn_set_crc_value", return_value=True) as second_set_crc:
                result = tab.fn_set_crc_value("47")

        self.assertTrue(result)
        first_set_crc.assert_not_called()
        second_set_crc.assert_called_once_with("47")

        tab.workspace_sidebar.select_workspace(first_workspace.id)
        with patch.object(first_panel, "fn_set_crc_value", return_value=True) as first_set_crc:
            result = tab.fn_set_crc_value("99")

        self.assertTrue(result)
        first_set_crc.assert_called_once_with("99")

    def test_coding_value_tab_crc_routing_returns_false_without_active_panel(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)
        workspace = tab.workspace_sidebar.workspaces()[0]

        tab.workspace_sidebar.remove_workspace(workspace.id)

        self.assertIsNone(tab.active_coding_value_panel())
        self.assertFalse(tab.fn_set_crc_value("47"))

    def test_coding_value_tab_theme_refresh_updates_sidebar_and_all_panels(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)
        first_workspace = tab.workspace_sidebar.workspaces()[0]
        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("ECU 2", True),
        ):
            tab.workspace_sidebar.btn_add.click()
        second_workspace = tab.workspace_sidebar.active_workspace()
        first_panel = tab._panel_by_workspace_id[first_workspace.id]
        second_panel = tab._panel_by_workspace_id[second_workspace.id]

        with patch.object(tab.workspace_sidebar, "fn_refresh_theme") as sidebar_theme:
            with patch.object(first_panel, "fn_refresh_theme") as first_panel_theme:
                with patch.object(second_panel, "fn_refresh_theme") as second_panel_theme:
                    tab.fn_refresh_theme()

        sidebar_theme.assert_called_once_with()
        first_panel_theme.assert_called_once_with()
        second_panel_theme.assert_called_once_with()

    def test_delete_workspace_removes_panel_and_keeps_active_fallback(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)
        first_workspace = tab.workspace_sidebar.workspaces()[0]
        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("ECU 2", True),
        ):
            tab.workspace_sidebar.btn_add.click()
        second_workspace = tab.workspace_sidebar.active_workspace()
        first_panel = tab._panel_by_workspace_id[first_workspace.id]
        second_panel = tab._panel_by_workspace_id[second_workspace.id]

        self.assertTrue(tab.workspace_sidebar.remove_workspace(second_workspace.id))

        self.assertEqual(tab.workspace_stack.count(), 1)
        self.assertNotIn(second_workspace.id, tab._panel_by_workspace_id)
        self.assertEqual(tab.workspace_stack.indexOf(second_panel), -1)
        self.assertIs(tab.workspace_sidebar.active_workspace(), first_workspace)
        self.assertIs(tab.workspace_stack.currentWidget(), first_panel)

    def test_delete_last_workspace_and_add_after_empty_is_safe(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)
        workspace = tab.workspace_sidebar.workspaces()[0]

        self.assertTrue(tab.workspace_sidebar.remove_workspace(workspace.id))

        self.assertEqual(tab.workspace_stack.count(), 0)
        self.assertEqual(tab._panel_by_workspace_id, {})
        self.assertIsNone(tab.active_coding_value_panel())

        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("ECU New", True),
        ):
            tab.workspace_sidebar.btn_add.click()

        self.assertEqual(tab.workspace_stack.count(), 1)
        self.assertEqual(len(tab._panel_by_workspace_id), 1)
        self.assertIs(tab.workspace_stack.currentWidget(), tab.active_coding_value_panel())

    def test_switching_ecu_workspaces_preserves_independent_panel_state(self):
        tab = self.create_tab()
        self.addCleanup(tab.deleteLater)

        def row(parameter, byte_pos):
            return CodingValueRow(
                parameter=parameter,
                byte_pos=str(byte_pos),
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            )

        mhu_workspace = tab.workspace_sidebar.workspaces()[0]
        tab.workspace_sidebar.rename_workspace(mhu_workspace.id, "MHU")
        with patch(
            "gui.widgets.coding_value.ecu_workspace_sidebar.QInputDialog.getText",
            return_value=("VCU", True),
        ):
            tab.workspace_sidebar.btn_add.click()
        vcu_workspace = tab.workspace_sidebar.active_workspace()

        mhu_definition = CodingDefinition(
            name="MHU Definition",
            rows=(row("MHU Byte 0", 0), row("MHU Byte 1", 1)),
        )
        vcu_definition = CodingDefinition(
            name="VCU Definition",
            rows=(row("VCU Byte 0", 0), row("VCU Byte 1", 1)),
        )

        tab.workspace_sidebar.select_workspace(mhu_workspace.id)
        mhu_panel = tab.active_coding_value_panel()
        mhu_panel._set_coding_definition(mhu_definition)
        mhu_panel.txt_coding_value.setPlainText("AA 11")
        mhu_panel.encode_coding_payload()
        mhu_panel.table.fn_set_raw_value_at_row(0, "AB")
        mhu_panel.check_coding_value()
        mhu_panel.filter_no_match_rows()

        tab.workspace_sidebar.select_workspace(vcu_workspace.id)
        vcu_panel = tab.active_coding_value_panel()
        vcu_panel._set_coding_definition(vcu_definition)
        vcu_panel.txt_coding_value.setPlainText("BB CC")
        vcu_panel.encode_coding_payload()
        vcu_panel.table.fn_set_raw_value_at_row(1, "CD")
        vcu_panel.check_coding_value()

        self.assertIsNot(mhu_panel, vcu_panel)
        self.assertIs(tab.workspace_stack.currentWidget(), vcu_panel)
        self.assertIs(vcu_panel._coding_definition, vcu_definition)
        self.assertEqual(vcu_panel.txt_coding_value.toPlainText(), "BB CC")
        self.assertEqual(vcu_panel.txt_coding_preview.toPlainText(), "BB CD")
        self.assertEqual(vcu_panel.table.item(0, 0).text(), "VCU Byte 0")
        self.assertEqual(vcu_panel.table.item(1, 4).text(), "CD")
        self.assertEqual(vcu_panel.txt_parameter_filter.text(), "")
        self.assertIn("Total No-M", vcu_panel.txt_working_log.toPlainText())

        tab.workspace_sidebar.select_workspace(mhu_workspace.id)

        self.assertIs(tab.workspace_stack.currentWidget(), mhu_panel)
        self.assertIs(mhu_panel._coding_definition, mhu_definition)
        self.assertEqual(mhu_panel.txt_coding_value.toPlainText(), "AA 11")
        self.assertEqual(mhu_panel.txt_coding_preview.toPlainText(), "AB 11")
        self.assertEqual(mhu_panel.table.item(0, 0).text(), "MHU Byte 0")
        self.assertEqual(mhu_panel.table.item(0, 4).text(), "AB")
        self.assertEqual(mhu_panel.txt_parameter_filter.text(), "No-M")
        self.assertIn("Total No-M", mhu_panel.txt_working_log.toPlainText())

        tab.workspace_sidebar.select_workspace(vcu_workspace.id)

        self.assertEqual(vcu_panel.txt_coding_preview.toPlainText(), "BB CD")
        self.assertEqual(vcu_panel.table.item(1, 4).text(), "CD")

    def test_set_crc_value_updates_crc_parameter_row(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)
        panel.table.set_rows([
            CodingValueRow(
                parameter="Vehicle Name",
                byte_pos="13",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
            CodingValueRow(
                parameter="Payload CRC Byte",
                byte_pos="66",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(
                    CodingValueOption(raw_value="0x47", label="47"),
                ),
            ),
        ])
        panel._payload_preview_bytes = [0] * 67
        panel.txt_coding_preview.setPlainText(
            panel._format_payload_bytes(panel._payload_preview_bytes)
        )

        self.assertTrue(panel.fn_set_crc_value("47"))

        self.assertEqual(panel.table.item(1, 4).text(), "47")
        self.assertEqual(panel.table.cellWidget(1, 5).currentText(), "47")
        self.assertEqual(panel._payload_preview_bytes[66], 0x47)

    def test_set_crc_value_offsets_data_byte_pos_for_preview_header(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)
        panel.table.set_rows([
            CodingValueRow(
                parameter="Payload CRC Byte",
                byte_pos="63",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(
                    CodingValueOption(raw_value="0x47", label="47"),
                ),
            ),
        ])
        payload_bytes = [0] * 67
        payload_bytes[63] = 0xAA
        panel.set_payload_preview(
            panel._format_payload_bytes(payload_bytes)
        )

        self.assertTrue(panel.fn_set_crc_value("47"))

        self.assertEqual(panel.table.item(0, 4).text(), "47")
        self.assertEqual(panel._payload_preview_bytes[63], 0xAA)
        self.assertEqual(panel._payload_preview_bytes[66], 0x47)

    def test_export_coding_files_dir_lives_under_config(self):
        self.assertEqual(
            EXPORT_CODING_FILES_DIR,
            CONFIG_DIR / "export_coding_files",
        )

    def test_text_selection_is_cleared_when_coding_lineedit_loses_focus(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)

        panel.file_path.setText("C:/tmp/coding.xlsx")
        panel.file_path.selectAll()
        self.assertNotEqual(panel.file_path.selectedText(), "")

        QApplication.sendEvent(
            panel.file_path,
            QFocusEvent(QEvent.FocusOut),
        )

        self.assertEqual(panel.file_path.selectedText(), "")

    def test_text_selection_is_cleared_when_coding_editor_loses_focus(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)

        panel.txt_coding_value.setPlainText("62 F1 08")
        panel.txt_coding_value.selectAll()
        self.assertNotEqual(panel.txt_coding_value.textCursor().selectedText(), "")

        QApplication.sendEvent(
            panel.txt_coding_value,
            QFocusEvent(QEvent.FocusOut),
        )

        self.assertEqual(panel.txt_coding_value.textCursor().selectedText(), "")
    def test_file_path_starts_empty(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)

        self.assertEqual(panel.file_path.text(), "")

    def test_filter_buttons_start_disabled_without_table_data(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)

        self.assertFalse(panel.btn_filter_no_m.isEnabled())
        self.assertFalse(panel.btn_refresh_filter.isEnabled())

    def test_json_drop_list_starts_empty_between_browse_and_import(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)

        file_layout = panel.file_row.layout()
        self.assertEqual(file_layout.indexOf(panel.btn_browse), 1)
        self.assertEqual(file_layout.indexOf(panel.cmb_coding_json), 2)
        self.assertEqual(file_layout.indexOf(panel.btn_import), 3)
        self.assertEqual(file_layout.indexOf(panel.btn_default), 4)
        self.assertEqual(panel.btn_default.text(), "Default")
        self.assertFalse(panel.btn_default.isEnabled())
        self.assertEqual(panel.cmb_coding_json.width(), 200)
        self.assertEqual(panel.cmb_coding_json.maxVisibleItems(), 5)
        self.assertIn(
            "QScrollBar:vertical",
            panel.cmb_coding_json.view().verticalScrollBar().styleSheet(),
        )
        self.assertEqual(panel.cmb_coding_json.currentText(), "")

    def test_coding_value_shortcuts_click_expected_action_buttons(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)
        panel.table.set_rows([
            CodingValueRow(
                parameter="Payload Byte",
                byte_pos="2",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
        ])
        panel.txt_coding_value.setPlainText("62 F1 08")
        panel.encode_coding_payload()
        panel.btn_filter_no_m.setEnabled(True)

        expected = {
            "Ctrl+R": panel.btn_check,
            "Ctrl+F": panel.btn_filter_no_m,
            "Ctrl+Z": panel.btn_refresh_filter,
            "Ctrl+Q": panel.btn_encode,
        }
        clicked = []
        for button in expected.values():
            button.clicked.connect(
                lambda checked=False, button=button: clicked.append(button)
            )

        shortcuts = {
            shortcut.key().toString(QKeySequence.PortableText): shortcut
            for shortcut in panel._coding_value_shortcuts
        }

        self.assertEqual(set(shortcuts), set(expected))
        for shortcut in shortcuts.values():
            self.assertIs(shortcut.parent(), panel)
            self.assertEqual(shortcut.context(), Qt.WidgetWithChildrenShortcut)
        for sequence, button in expected.items():
            shortcuts[sequence].activated.emit()
            self.assertIs(clicked[-1], button)

    def test_import_coding_excel_renders_table_with_decoded_combobox(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "DID Number (hex)",
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "0xF112",
            "Vehicle Name",
            3,
            0,
            8,
            "0x1=VF3\n0xFF=Invalid",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("62 F1 12 FF")
            panel.import_coding_value()

        self.assertEqual(panel.table.rowCount(), 1)
        self.assertEqual(panel.table.item(0, 0).text(), "Vehicle Name")
        self.assertEqual(panel.table.item(0, 4).text(), "")

        combo = panel.table.cellWidget(0, 5)
        self.assertEqual(combo.currentText(), "")
        self.assertEqual(combo.count(), 2)

        panel.encode_coding_payload()
        self.assertEqual(panel.table.item(0, 4).text(), "FF")
        self.assertEqual(combo.currentText(), "Invalid")

        combo.setCurrentIndex(combo.findText("VF3"))
        self.assertEqual(panel.table.item(0, 4).text(), "01")

        panel.table.item(0, 4).setText("FF")
        self.assertEqual(combo.currentText(), "Invalid")
        panel.table.item(0, 4).setText("01")
        self.assertEqual(combo.currentText(), "VF3")

    def test_default_button_fills_zero_payload_matching_imported_file_length(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            13,
            0,
            8,
            "0x00=Default\n0x09=VF7NP",
        ])
        sheet.append([
            "Two Byte Value",
            14,
            0,
            16,
            "0x0000=Default",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.import_coding_value()

        self.assertTrue(panel.btn_default.isEnabled())

        panel.btn_default.click()

        self.assertEqual(
            panel.txt_coding_value.toPlainText(),
            " ".join(["00"] * 16),
        )

    def test_import_coding_excel_exports_parsed_rows_to_json(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            13,
            0,
            8,
            "0x03=VF3\n0x09=VF7NP",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            export_dir = Path(tmpdir) / "config" / "export_coding_files"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))

            with patch(
                "gui.widgets.coding_value.coding_value_panel.EXPORT_CODING_FILES_DIR",
                export_dir,
            ):
                panel.import_coding_value()

            exported_path = export_dir / "coding.json"
            self.assertTrue(exported_path.exists())
            exported = json.loads(exported_path.read_text(encoding="utf-8"))

        self.assertEqual(exported["source_file"], "coding.xlsx")
        self.assertEqual(exported["rows"][0]["parameter"], "Vehicle Name")
        self.assertEqual(exported["rows"][0]["byte_pos"], "13")
        self.assertEqual(exported["rows"][0]["bit_pos"], "0")
        self.assertEqual(exported["rows"][0]["bit_length"], "8")
        self.assertEqual(
            exported["rows"][0]["decoded_options"],
            [
                {"raw_value": "0x03", "label": "VF3"},
                {"raw_value": "0x09", "label": "VF7NP"},
            ],
        )
    def test_selecting_json_drop_list_loads_table_from_export_folder(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            export_dir = Path(tmpdir) / "config" / "export_coding_files"
            export_dir.mkdir(parents=True)
            json_path = export_dir / "existing.json"
            json_path.write_text(
                json.dumps({
                    "source_file": "existing.xlsx",
                    "rows": [
                        {
                            "parameter": "Method Type",
                            "byte_pos": "14",
                            "bit_pos": "0",
                            "bit_length": "8",
                            "raw_value": "",
                            "decoded_value": "",
                            "decoded_options": [
                                {"raw_value": "0x05", "label": "Sky"},
                            ],
                        },
                    ],
                }),
                encoding="utf-8",
            )

            with patch(
                "gui.widgets.coding_value.coding_value_panel.EXPORT_CODING_FILES_DIR",
                export_dir,
            ):
                panel = CodingValuePanel()
                self.addCleanup(panel.deleteLater)
                panel.cmb_coding_json.setCurrentIndex(
                    panel.cmb_coding_json.findText("existing")
                )

        self.assertEqual(panel._coding_definition.name, "existing")
        self.assertEqual(panel._coding_definition.source_file, "existing.xlsx")
        self.assertEqual(panel.table.rowCount(), 1)
        self.assertEqual(panel.table.item(0, 0).text(), "Method Type")
        self.assertEqual(panel.table.item(0, 4).text(), "")
        self.assertEqual(panel.table.cellWidget(0, 5).currentText(), "")
    def test_import_button_loads_selected_json_without_excel_warning(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            export_dir = Path(tmpdir) / "config" / "export_coding_files"
            export_dir.mkdir(parents=True)
            json_path = export_dir / "existing.json"
            json_path.write_text(
                json.dumps({
                    "source_file": "existing.xlsx",
                    "rows": [
                        {
                            "parameter": "Method Type",
                            "byte_pos": "14",
                            "bit_pos": "0",
                            "bit_length": "8",
                            "raw_value": "",
                            "decoded_value": "",
                            "decoded_options": [
                                {"raw_value": "0x05", "label": "Sky"},
                            ],
                        },
                    ],
                }),
                encoding="utf-8",
            )

            with patch(
                "gui.widgets.coding_value.coding_value_panel.EXPORT_CODING_FILES_DIR",
                export_dir,
            ):
                panel = CodingValuePanel()
                self.addCleanup(panel.deleteLater)
                panel.cmb_coding_json.setCurrentIndex(
                    panel.cmb_coding_json.findText("existing")
                )
                panel.file_path.clear()

                with patch(
                    "gui.widgets.coding_value.coding_value_panel.QMessageBox.warning"
                ) as warning:
                    panel.import_coding_value()

        warning.assert_not_called()
        self.assertEqual(panel._coding_definition.name, "existing")
        self.assertEqual(panel.table.rowCount(), 1)
        self.assertEqual(panel.table.item(0, 0).text(), "Method Type")
    def test_import_coding_excel_refreshes_json_drop_list_and_loads_table_from_json(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            13,
            0,
            8,
            "0x03=VF3\n0x09=VF7NP",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            export_dir = Path(tmpdir) / "config" / "export_coding_files"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))

            with patch(
                "gui.widgets.coding_value.coding_value_panel.EXPORT_CODING_FILES_DIR",
                export_dir,
            ):
                panel.import_coding_value()

        self.assertEqual(panel.cmb_coding_json.currentText(), "coding")
        self.assertEqual(panel.file_path.text(), "")
        self.assertEqual(panel._coding_definition.name, "coding")
        self.assertEqual(panel._coding_definition.rows[0].parameter, "Vehicle Name")
        self.assertEqual(panel.table.rowCount(), 1)
        self.assertEqual(panel.table.item(0, 0).text(), "Vehicle Name")
        self.assertEqual(panel.table.item(0, 4).text(), "")
        self.assertEqual(panel.table.cellWidget(0, 5).currentText(), "")
    def test_decoded_combobox_options_are_one_item_per_method_type_line(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "$MethodType",
        ])
        sheet.append([
            "Vehicle Variant",
            14,
            0,
            8,
            "\n0x0=Unsupported\n0x1= Comfort\n0x2= Premium\n"
            "0x3= Earth\n0x4= Wind\n0x5=  Sky\n"
            "0x6=  HerioGreen\n0x7=  Reserved\n"
            "0x8 = Reserved\n0x9 = Reserved\n"
            "0xA = Reserved\n0xF=Invalid",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("00 " * 14 + "02")
            panel.import_coding_value()

        combo = panel.table.cellWidget(0, 5)
        self.assertEqual(combo.currentText(), "")
        panel.encode_coding_payload()
        self.assertEqual(combo.currentText(), "Premium")
        self.assertEqual(combo.count(), 12)
        self.assertFalse(combo.isEditable())
        self.assertEqual(combo.maxVisibleItems(), 5)
        self.assertEqual(combo.itemText(1), "Comfort")
        self.assertEqual(combo.itemText(2), "Premium")
        self.assertIn(
            "QScrollBar:vertical",
            combo.view().verticalScrollBar().styleSheet(),
        )
        self.assertGreaterEqual(
            panel.table.rowHeight(0),
            combo.sizeHint().height(),
        )

    def test_raw_value_edit_formats_by_bit_length_and_rejects_overflow(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Two Byte Value",
            13,
            0,
            16,
            "0x123=Small\n0x1234=Large",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("00 " * 13 + "12 34")
            panel.import_coding_value()
            panel.encode_coding_payload()

        raw_item = panel.table.item(0, 4)
        combo = panel.table.cellWidget(0, 5)

        self.assertTrue(raw_item.flags() & Qt.ItemIsEditable)
        self.assertEqual(raw_item.text(), "12 34")
        self.assertEqual(combo.currentText(), "Large")

        raw_item.setText("123")
        self.assertEqual(raw_item.text(), "01 23")
        self.assertEqual(combo.currentText(), "Small")

        with patch(
            "gui.widgets.coding_value.coding_value_table.QMessageBox.warning"
        ) as warning:
            raw_item.setText("1 00 00")

        warning.assert_called_once()
        self.assertEqual(raw_item.text(), "01 23")
        self.assertEqual(combo.currentText(), "Small")

    def test_raw_value_keeps_two_hex_digits_for_nibble_values(self):
        self.assertEqual(
            CodingValueTable._normalize_user_raw_value("03", 4),
            "03",
        )
        self.assertEqual(
            CodingValueTable._normalize_user_raw_value("1", 4),
            "01",
        )
        self.assertEqual(
            CodingValueTable._normalize_user_raw_value("D", 4),
            "0D",
        )
        self.assertIsNone(
            CodingValueTable._normalize_user_raw_value("10", 4),
        )

    def test_raw_value_accepts_multiple_nibble_tokens_when_bit_length_is_text(self):
        self.assertEqual(
            CodingValueTable._normalize_user_raw_value("08", "4.0"),
            "08",
        )
        self.assertEqual(
            CodingValueTable._normalize_user_raw_value("08 09", "4.0"),
            "08 09",
        )
        self.assertEqual(
            CodingValueTable._normalize_user_raw_value("8 9", "4 bit"),
            "08 09",
        )
        self.assertIsNone(
            CodingValueTable._normalize_user_raw_value("10 09", "4 bit"),
        )

    def test_encode_payload_buttons_update_raw_values_by_byte_position(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            3,
            0,
            8,
            "0x37=Seven\n0xFF=Invalid",
        ])
        sheet.append([
            "Two Byte Value",
            4,
            0,
            16,
            "0x4A39=Name",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("62 F1 08 00 00 00")
            panel.import_coding_value()
            panel.txt_coding_value.setPlainText("62 F1 08 37 4A 39")
            panel.encode_coding_payload()

        self.assertEqual(panel.table.item(0, 4).text(), "37")
        self.assertEqual(panel.table.cellWidget(0, 5).currentText(), "Seven")
        self.assertEqual(panel.table.item(1, 4).text(), "4A 39")
        self.assertEqual(panel.table.cellWidget(1, 5).currentText(), "Name")

    def test_encode_payload_splits_bit_fields_by_bit_position(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Byte 17 Low Nibble",
            17,
            0,
            4,
            "0x2=Low",
        ])
        sheet.append([
            "Byte 17 High Nibble",
            17,
            4,
            4,
            "0x1=High",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("00 " * 18)
            panel.import_coding_value()
            panel.txt_coding_value.setPlainText("00 " * 17 + "12")
            panel.encode_coding_payload()

        self.assertEqual(panel.table.item(0, 4).text(), "02")
        self.assertEqual(panel.table.cellWidget(0, 5).currentText(), "Low")
        self.assertEqual(panel.table.item(1, 4).text(), "01")
        self.assertEqual(panel.table.cellWidget(1, 5).currentText(), "High")

    def test_payload_copy_and_clear_buttons(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)
        panel.txt_coding_value.setPlainText("62 F1 08")

        panel.copy_coding_payload()

        self.assertEqual(
            QApplication.clipboard().text(),
            "62 F1 08",
        )

        panel.txt_coding_preview.setPlainText("62 F1 99")
        QApplication.clipboard().setText("keep clipboard")

        panel.clear_coding_payload()

        self.assertEqual(panel.txt_coding_value.toPlainText(), "")
        self.assertEqual(panel.txt_coding_preview.toPlainText(), "")
        self.assertEqual(QApplication.clipboard().text(), "keep clipboard")

    def test_payload_input_and_actions_share_one_compact_row(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)

        payload_layout = panel.payload_row.layout()
        input_layout = panel.payload_input_row.layout()
        preview_layout = panel.payload_preview_row.layout()

        self.assertIsInstance(payload_layout, QVBoxLayout)
        self.assertIsInstance(input_layout, QHBoxLayout)
        self.assertIsInstance(preview_layout, QHBoxLayout)
        self.assertEqual(input_layout.indexOf(panel.btn_encode), 0)
        self.assertEqual(input_layout.indexOf(panel.btn_copy), 1)
        self.assertEqual(input_layout.indexOf(panel.btn_clear), 2)
        self.assertEqual(input_layout.indexOf(panel.txt_coding_value), 3)
        self.assertEqual(preview_layout.indexOf(panel.btn_preview_refresh), 0)
        self.assertEqual(preview_layout.indexOf(panel.btn_copy_preview), 1)
        self.assertEqual(preview_layout.indexOf(panel.btn_preview_edit), 2)
        self.assertEqual(preview_layout.indexOf(panel.txt_coding_preview), 3)
        self.assertLessEqual(
            panel.txt_coding_value.maximumHeight(),
            panel.txt_coding_value.fontMetrics().lineSpacing() * 3 + 24,
        )
        self.assertIn(
            "QScrollBar:vertical",
            panel.txt_coding_value.verticalScrollBar().styleSheet(),
        )
        self.assertLessEqual(
            panel.txt_coding_preview.maximumHeight(),
            panel.txt_coding_preview.fontMetrics().lineSpacing() * 3 + 24,
        )
        self.assertIn(
            "QScrollBar:vertical",
            panel.txt_coding_preview.verticalScrollBar().styleSheet(),
        )

    def test_action_buttons_require_payload_data_after_import(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Payload Byte",
            2,
            0,
            8,
            "0x08=Old",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.import_coding_value()

        self.assertFalse(panel.btn_check.isEnabled())
        self.assertFalse(panel.btn_copy_data_payload.isEnabled())
        self.assertFalse(panel.btn_export.isEnabled())
        self.assertFalse(panel.btn_table_clear.isEnabled())
        self.assertFalse(panel.btn_preview_edit.isEnabled())
        self.assertFalse(panel.btn_copy_preview.isEnabled())
        self.assertFalse(panel.btn_clear.isEnabled())
        self.assertFalse(panel.btn_filter_no_m.isEnabled())
        self.assertTrue(panel.btn_refresh_filter.isEnabled())
        self.assertEqual(panel.table.columnCount(), 6)

        panel.check_coding_value()
        self.assertEqual(panel.table.columnCount(), 6)

        panel.txt_coding_value.setPlainText("62 F1 08")
        self.assertTrue(panel.btn_encode.isEnabled())
        self.assertTrue(panel.btn_copy.isEnabled())
        self.assertTrue(panel.btn_clear.isEnabled())
        self.assertFalse(panel.btn_check.isEnabled())

        panel.encode_coding_payload()
        self.assertTrue(panel.btn_check.isEnabled())
        self.assertFalse(panel.btn_copy_data_payload.isEnabled())
        self.assertTrue(panel.btn_export.isEnabled())
        self.assertTrue(panel.btn_table_clear.isEnabled())
        self.assertFalse(panel.btn_filter_no_m.isEnabled())
        self.assertTrue(panel.btn_refresh_filter.isEnabled())
        self.assertFalse(panel.btn_preview_edit.isEnabled())
        self.assertTrue(panel.btn_copy_preview.isEnabled())

        panel.check_coding_value()
        self.assertTrue(panel.btn_filter_no_m.isEnabled())
        self.assertTrue(panel.btn_refresh_filter.isEnabled())
        self.assertTrue(panel.btn_preview_edit.isEnabled())

        panel.clear_coding_payload()
        self.assertFalse(panel.btn_check.isEnabled())
        self.assertFalse(panel.btn_copy_data_payload.isEnabled())
        self.assertFalse(panel.btn_export.isEnabled())
        self.assertFalse(panel.btn_table_clear.isEnabled())
        self.assertFalse(panel.btn_preview_edit.isEnabled())
        self.assertFalse(panel.btn_copy_preview.isEnabled())
        self.assertFalse(panel.btn_clear.isEnabled())
        self.assertFalse(panel.btn_filter_no_m.isEnabled())
        self.assertTrue(panel.btn_refresh_filter.isEnabled())

    def test_check_button_shows_fixed_before_columns_from_encoded_baseline(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Encoded Byte",
            2,
            0,
            8,
            "0x08=Old\n0x99=New",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.import_coding_value()

        panel.txt_coding_value.setPlainText("62 F1 08")
        panel.encode_coding_payload()
        panel.table.item(0, 4).setText("99")
        panel.check_coding_value()

        self.assertEqual(panel.table.item(0, 6).text(), "08")
        self.assertEqual(panel.table.item(0, 7).text(), "Old")
        self.assertEqual(panel.table.item(0, 8).text(), "No-M")

    def test_checked_result_updates_but_before_columns_stay_fixed(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Editable Byte",
            2,
            0,
            8,
            "0x08=Old\n0x99=New",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("62 F1 08")
            panel.import_coding_value()
            panel.encode_coding_payload()

        panel.check_coding_value()
        combo = panel.table.cellWidget(0, 5)
        combo.setCurrentIndex(combo.findText("New"))

        self.assertEqual(panel.table.item(0, 4).text(), "99")
        self.assertEqual(panel.table.item(0, 6).text(), "08")
        self.assertEqual(panel.table.item(0, 7).text(), "Old")
        self.assertEqual(panel.table.item(0, 8).text(), "No-M")

    def test_check_button_writes_working_log_for_no_match_changes_and_crc(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            13,
            0,
            8,
            "0x09=VF7NP\n0x01=VF5",
        ])
        sheet.append([
            "Method Type",
            14,
            0,
            8,
            "0x06=Heriogreen\n0x05=Sky",
        ])
        sheet.append([
            "Payload CRC Byte",
            66,
            0,
            8,
            "0xA5=A5\n0x12=12",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText(
                "00 " * 13 + "09 06 " + "00 " * 51 + "A5"
            )
            panel.import_coding_value()
            panel.encode_coding_payload()

        panel.table.cellWidget(0, 5).setCurrentIndex(
            panel.table.cellWidget(0, 5).findText("VF5")
        )
        panel.table.cellWidget(1, 5).setCurrentIndex(
            panel.table.cellWidget(1, 5).findText("Sky")
        )
        panel.table.cellWidget(2, 5).setCurrentIndex(
            panel.table.cellWidget(2, 5).findText("12")
        )
        panel._payload_baseline_bytes[-1] = 0x77
        panel._payload_preview_bytes[-1] = 0x88
        panel.check_coding_value()

        self.assertEqual(
            panel.txt_working_log.toPlainText(),
            "Total No-M: 3\n\n"
            "Changes: Decode value before -> after\n\n"
            "Byte 13: VF7NP -> VF5\n"
            "Byte 14: Heriogreen -> Sky\n"
            "Byte 66: A5 -> 12\n\n"
            "CRC: 77 -> 88",
        )

    def test_check_button_updates_crc_after_short_vin(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)
        panel.table.set_rows([
            CodingValueRow(
                parameter="Short VIN",
                byte_pos="2",
                bit_pos="0",
                bit_length="16",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
            CodingValueRow(
                parameter="Payload CRC Byte",
                byte_pos="6",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
        ])
        payload = [0x10, 0x20, 0xAA, 0xBB, 0x01, 0x02, 0x00]
        panel.txt_coding_value.setPlainText(panel._format_payload_bytes(payload))
        panel.encode_coding_payload()
        expected_crc = calculate_crc8_sae_j1850(bytes([0x01, 0x02]))

        panel.check_coding_value()

        self.assertEqual(panel.table.item(1, 4).text(), f"{expected_crc:02X}")
        self.assertEqual(panel._payload_preview_bytes[6], expected_crc)
        self.assertEqual(
            panel.txt_coding_preview.toPlainText(),
            panel._format_payload_bytes(payload[:-1] + [expected_crc]),
        )

    def test_check_button_updates_crc_after_format_byte_when_short_vin_is_missing(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)
        panel.table.set_rows([
            CodingValueRow(
                parameter="Format Byte",
                byte_pos="3",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
            CodingValueRow(
                parameter="Payload CRC Byte",
                byte_pos="6",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
        ])
        payload = [0x62, 0xF1, 0x08, 0x99, 0x01, 0x02, 0x00]
        panel.txt_coding_value.setPlainText(panel._format_payload_bytes(payload))
        panel.encode_coding_payload()
        expected_crc = calculate_crc8_sae_j1850(bytes([0x01, 0x02]))

        panel.check_coding_value()

        self.assertEqual(panel.table.item(1, 4).text(), f"{expected_crc:02X}")
        self.assertEqual(panel._payload_preview_bytes[6], expected_crc)
        self.assertEqual(
            panel.txt_coding_preview.toPlainText(),
            panel._format_payload_bytes(payload[:-1] + [expected_crc]),
        )

    def test_check_button_updates_crc_from_data_payload_offsets_with_uds_header(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)
        panel.table.set_rows([
            CodingValueRow(
                parameter="Format Byte",
                byte_pos="0",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
            CodingValueRow(
                parameter="Payload CRC Byte",
                byte_pos="3",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
        ])
        payload = [0x62, 0xF1, 0x08, 0x99, 0x01, 0x02, 0x00]
        panel.txt_coding_value.setPlainText(panel._format_payload_bytes(payload))
        panel.encode_coding_payload()
        expected_crc = calculate_crc8_sae_j1850(bytes([0x01, 0x02]))

        panel.check_coding_value()

        self.assertEqual(panel.table.item(1, 4).text(), f"{expected_crc:02X}")
        self.assertEqual(panel._payload_preview_bytes[6], expected_crc)
        self.assertEqual(
            panel.txt_coding_preview.toPlainText(),
            panel._format_payload_bytes(payload[:-1] + [expected_crc]),
        )

    def test_parameter_filter_hides_non_matching_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            13,
            0,
            8,
            "0x03=VF3",
        ])
        sheet.append([
            "Method Type",
            14,
            0,
            8,
            "0x06=Heriogreen",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.import_coding_value()

        self.assertEqual(panel.table.rowCount(), 2)
        filter_layout = panel.filter_row.layout()
        self.assertEqual(filter_layout.indexOf(panel.txt_parameter_filter), 0)
        self.assertEqual(filter_layout.indexOf(panel.filter_action_row), 1)
        self.assertEqual(panel.filter_action_row.width(), 220)
        self.assertEqual(
            panel.btn_filter_no_m.width(),
            panel.btn_refresh_filter.width(),
        )
        self.assertEqual(panel.btn_filter_no_m.width(), 106)

        panel.table.selectRow(1)
        panel.txt_parameter_filter.setText("vehicle")

        self.assertFalse(panel.table.isRowHidden(0))
        self.assertTrue(panel.table.isRowHidden(1))
        self.assertEqual(panel.table.selectedItems(), [])

        panel.txt_parameter_filter.setText("")

        self.assertFalse(panel.table.isRowHidden(0))
        self.assertFalse(panel.table.isRowHidden(1))
    def test_filter_no_match_and_refresh_buttons_update_filter(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Unchanged Byte",
            2,
            0,
            8,
            "0x08=Old",
        ])
        sheet.append([
            "Changed Byte",
            3,
            0,
            8,
            "0x37=Old\n0x99=New",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("62 F1 08 37")
            panel.import_coding_value()
            panel.encode_coding_payload()

        self.assertFalse(panel.btn_filter_no_m.isEnabled())
        self.assertTrue(panel.btn_refresh_filter.isEnabled())

        panel.check_coding_value()
        panel.table.item(1, 4).setText("99")

        panel.btn_filter_no_m.click()

        self.assertEqual(panel.txt_parameter_filter.text(), "No-M")
        self.assertTrue(panel.table.isRowHidden(0))
        self.assertFalse(panel.table.isRowHidden(1))

        panel.btn_refresh_filter.click()

        self.assertEqual(panel.txt_parameter_filter.text(), "")
        self.assertFalse(panel.table.isRowHidden(0))
        self.assertFalse(panel.table.isRowHidden(1))

    def test_export_button_writes_visible_table_to_coding_value_report(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            13,
            0,
            8,
            "0x03=VF3\n0x09=VF7NP",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            coding_path = Path(tmpdir) / "coding.xlsx"
            report_root = Path(tmpdir) / "report"
            custom_export_path = Path(tmpdir) / "custom_coding_report.xlsx"
            workbook.save(coding_path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(coding_path))
            panel.txt_coding_value.setPlainText(
                "62 F1 08 00 37 4A 39 39 39 37 34 31 26 03"
            )
            panel.import_coding_value()
            panel.encode_coding_payload()
            panel.check_coding_value()

            panel._report_service = CodingValueReportService(report_root=report_root)
            with patch(
                "gui.widgets.coding_value.coding_value_panel.QFileDialog.getSaveFileName",
                return_value=(str(custom_export_path), "Excel Files (*.xlsx)"),
            ) as save_dialog, patch(
                "gui.widgets.coding_value.coding_value_panel.QMessageBox.information"
            ) as information:
                panel.btn_export.click()

            save_dialog.assert_called_once()
            default_path = Path(save_dialog.call_args.args[2])
            self.assertEqual(default_path.parent.parent, report_root)
            self.assertTrue(default_path.name.startswith("EEIV_report_Coding value_"))
            information.assert_called_once()
            exported_path = Path(information.call_args.args[2].split("\n\n")[1])
            self.assertEqual(exported_path, custom_export_path)
            self.assertTrue(exported_path.exists())
            self.assertEqual(exported_path.suffix, ".xlsx")

            exported = load_workbook(exported_path)
            self.assertEqual(exported.sheetnames, ["Coding value_62 F1 08"])
            exported_sheet = exported["Coding value_62 F1 08"]
            self.assertEqual(
                [cell.value for cell in exported_sheet[1]],
                [
                    "Parameter",
                    "Byte Pos",
                    "Bit Pos",
                    "Bit Lengh",
                    "Raw value",
                    "Decoded Value (Editable)",
                    "Raw value (before)",
                    "Decoded value (before)",
                    "Result",
                ],
            )
            self.assertEqual(exported_sheet.cell(row=2, column=1).value, "Vehicle Name")
            self.assertEqual(exported_sheet.cell(row=2, column=5).value, "03")
            self.assertEqual(exported_sheet.cell(row=2, column=6).value, "VF3")
            self.assertEqual(exported_sheet.cell(row=2, column=9).value, "MATCH")
            self.assertEqual(exported_sheet.freeze_panes, "A2")
            self.assertEqual(exported_sheet.auto_filter.ref, exported_sheet.dimensions)
            self.assertTrue(exported_sheet["A1"].font.bold)
            self.assertTrue(_excel_color(exported_sheet["A1"]).endswith(COLOR_HEADER))
            self.assertEqual(exported_sheet["A1"].alignment.horizontal, "center")
            self.assertEqual(exported_sheet["A1"].border.left.style, "thin")
            self.assertEqual(exported_sheet["A2"].border.left.style, "thin")
    def test_export_button_highlights_no_match_rows_with_warning_fill(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            13,
            0,
            8,
            "0x03=VF3\n0x09=VF7NP",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            coding_path = Path(tmpdir) / "coding.xlsx"
            report_root = Path(tmpdir) / "report"
            custom_export_path = Path(tmpdir) / "custom_coding_report.xlsx"
            workbook.save(coding_path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(coding_path))
            panel.txt_coding_value.setPlainText(
                "62 F1 08 00 37 4A 39 39 39 37 34 31 26 03"
            )
            panel.import_coding_value()
            panel.encode_coding_payload()
            panel.table.cellWidget(0, 5).setCurrentIndex(
                panel.table.cellWidget(0, 5).findText("VF7NP")
            )
            panel.check_coding_value()

            panel._report_service = CodingValueReportService(report_root=report_root)
            with patch(
                "gui.widgets.coding_value.coding_value_panel.QFileDialog.getSaveFileName",
                return_value=(str(custom_export_path), "Excel Files (*.xlsx)"),
            ) as save_dialog, patch(
                "gui.widgets.coding_value.coding_value_panel.QMessageBox.information"
            ) as information:
                panel.btn_export.click()

            exported_path = Path(information.call_args.args[2].split("\n\n")[1])
            exported = load_workbook(exported_path)
            exported_sheet = exported["Coding value_62 F1 08"]
            self.assertEqual(exported_sheet.cell(row=2, column=9).value, "No-M")
            for cell in exported_sheet[2]:
                self.assertEqual(cell.fill.fill_type, "solid")
                self.assertTrue(cell.fill.fgColor.rgb.upper().endswith(COLOR_WARNING))
                self.assertEqual(cell.border.left.style, "thin")

    def test_table_clear_button_clears_table_and_working_log(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            13,
            0,
            8,
            "0x09=VF7NP\n0x01=VF5",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("00 " * 13 + "09")
            panel.import_coding_value()
            panel.encode_coding_payload()

        panel.table.cellWidget(0, 5).setCurrentIndex(
            panel.table.cellWidget(0, 5).findText("VF5")
        )
        panel.check_coding_value()
        self.assertEqual(panel.table.item(0, 4).text(), "01")
        self.assertEqual(panel.table.item(0, 6).text(), "09")
        self.assertEqual(panel.table.item(0, 7).text(), "VF7NP")
        self.assertEqual(panel.table.item(0, 8).text(), "No-M")
        self.assertNotEqual(panel.txt_working_log.toPlainText(), "")

        panel.btn_table_clear.click()

        self.assertEqual(panel.table.rowCount(), 0)
        self.assertEqual(panel.txt_working_log.toPlainText(), "")
        self.assertEqual(panel.txt_coding_preview.toPlainText(), "")
        self.assertFalse(panel.btn_check.isEnabled())
        self.assertFalse(panel.btn_table_clear.isEnabled())
    def test_check_button_clears_table_selection(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Selectable Byte",
            2,
            0,
            8,
            "0x08=Old\n0x99=New",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("62 F1 08")
            panel.import_coding_value()
            panel.encode_coding_payload()

        panel.table.selectRow(0)
        panel.table.setCurrentCell(0, 4)
        self.assertTrue(panel.table.selectedItems())
        self.assertEqual(panel.table.currentRow(), 0)

        panel.check_coding_value()

        self.assertEqual(panel.table.selectedItems(), [])
        self.assertEqual(panel.table.currentRow(), -1)
        self.assertEqual(panel.table.currentColumn(), -1)
    def test_check_button_adds_before_columns_and_compares_current_values(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Unchanged Byte",
            2,
            0,
            8,
            "0x08=Old\n0x99=New",
        ])
        sheet.append([
            "Changed Byte",
            3,
            0,
            8,
            "0x37=Seven\n0xFF=Invalid",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("62 F1 08 37")
            panel.import_coding_value()
            panel.encode_coding_payload()

        panel.check_coding_value()

        headers = [
            panel.table.horizontalHeaderItem(index).text()
            for index in range(panel.table.columnCount())
        ]
        self.assertEqual(headers[-3:], [
            "Raw value (before)",
            "Decoded value (before)",
            "Result",
        ])
        self.assertEqual(panel.table.item(0, 6).text(), "08")
        self.assertEqual(panel.table.item(0, 7).text(), "Old")
        self.assertEqual(panel.table.item(0, 8).text(), "MATCH")
        self.assertEqual(
            panel.table.item(0, 8).background().color().name().upper(),
            ThemeManager.fn_colors().SUCCESS.upper(),
        )
        self.assertEqual(panel.table.item(1, 6).text(), "37")
        self.assertEqual(panel.table.item(1, 7).text(), "Seven")
        self.assertEqual(panel.table.item(1, 8).text(), "MATCH")

        panel.table.item(1, 4).setText("FF")
        self.assertEqual(panel.table.item(1, 6).text(), "37")
        self.assertEqual(panel.table.item(1, 7).text(), "Seven")
        self.assertEqual(panel.table.item(1, 8).text(), "No-M")
        self.assertEqual(
            panel.table.item(1, 8).background().color().name().upper(),
            ThemeManager.fn_colors().WARNING.upper(),
        )

        panel.check_coding_value()
        self.assertEqual(panel.table.columnCount(), 9)
    def test_parameter_filter_matches_result_column_after_check(self):
        table = CodingValueTable()
        self.addCleanup(table.deleteLater)
        table.set_rows([
            CodingValueRow(
                parameter="Unchanged Byte",
                byte_pos="0",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
            CodingValueRow(
                parameter="Changed Byte",
                byte_pos="1",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
        ])
        table.encode_payload("01 02")
        table.item(1, 4).setText("FF")
        table.apply_check_results()

        table.filter_by_parameter("No-M")

        self.assertTrue(table.isRowHidden(0))
        self.assertFalse(table.isRowHidden(1))

        table.filter_by_parameter("")

        self.assertFalse(table.isRowHidden(0))
        self.assertFalse(table.isRowHidden(1))

    def test_result_header_click_does_not_reorder_table(self):
        table = CodingValueTable()
        self.addCleanup(table.deleteLater)
        table.set_rows([
            CodingValueRow(
                parameter="Original First",
                byte_pos="0",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
            CodingValueRow(
                parameter="Changed Second",
                byte_pos="1",
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            ),
        ])
        table.encode_payload("01 02")
        table.item(1, 4).setText("FF")
        table.apply_check_results()

        table.horizontalHeader().sectionClicked.emit(8)

        self.assertEqual(table.item(0, 0).text(), "Original First")
        self.assertEqual(table.item(1, 0).text(), "Changed Second")

    def test_table_uses_compact_widths_and_elides_overflow_text(self):
        table = CodingValueTable()
        self.addCleanup(table.deleteLater)

        self.assertEqual(table.textElideMode(), Qt.ElideRight)
        self.assertEqual(table.columnWidth(0), table.columnWidth(5))
        self.assertLess(table.columnWidth(1), table.columnWidth(0))
        self.assertLess(table.columnWidth(2), table.columnWidth(0))
        self.assertLess(table.columnWidth(3), table.columnWidth(0))
        self.assertLess(table.columnWidth(4), table.columnWidth(0))

        table.apply_check_results()

        self.assertGreater(table.columnWidth(7), table.columnWidth(1))
        self.assertLess(table.columnWidth(7), table.columnWidth(0))
        self.assertEqual(
            table.horizontalHeader().sectionResizeMode(7),
            QHeaderView.Stretch,
        )

    def test_copy_data_payload_button_copies_preview_after_first_three_bytes(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)

        panel.set_payload_preview("62 F1 08 01 0A ff")
        panel._update_action_states()

        self.assertTrue(panel.btn_copy_data_payload.isEnabled())

        panel.btn_copy_data_payload.click()

        self.assertEqual(
            QApplication.clipboard().text(),
            "01 0A FF",
        )

        panel.set_payload_preview("62 F1 08")
        panel._update_action_states()

        self.assertFalse(panel.btn_copy_data_payload.isEnabled())

    def test_table_has_right_side_actions_and_working_log_layout(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)

        table_area_layout = panel.table_area.layout()
        side_layout = panel.table_side_panel.layout()
        action_layout = panel.table_action_panel.layout()

        self.assertIsInstance(table_area_layout, QHBoxLayout)
        self.assertEqual(table_area_layout.indexOf(panel.table), 0)
        self.assertEqual(table_area_layout.indexOf(panel.table_side_panel), 1)
        self.assertIsInstance(side_layout, QVBoxLayout)
        self.assertEqual(side_layout.indexOf(panel.table_action_panel), 0)
        self.assertEqual(side_layout.indexOf(panel.txt_working_log), 1)
        self.assertIsInstance(action_layout, QVBoxLayout)
        self.assertEqual(action_layout.indexOf(panel.btn_check), 0)
        self.assertEqual(action_layout.indexOf(panel.btn_copy_data_payload), 1)
        self.assertEqual(action_layout.indexOf(panel.btn_export), 2)
        self.assertEqual(action_layout.indexOf(panel.btn_table_clear), 3)
        self.assertTrue(panel.txt_working_log.isReadOnly())
        self.assertIn(
            "QScrollBar:vertical",
            panel.txt_working_log.verticalScrollBar().styleSheet(),
        )
    def test_payload_preview_updates_and_highlights_table_raw_changes(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Vehicle Name",
            3,
            0,
            8,
            "0x37=Seven\n0xFF=Invalid",
        ])
        sheet.append([
            "Byte 4 Low Nibble",
            4,
            0,
            4,
            "0x2=Low\n0x5=Changed",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("62 F1 08 37 12")
            panel.import_coding_value()
            panel.encode_coding_payload()

        self.assertIsInstance(panel.txt_coding_preview, QPlainTextEdit)
        self.assertTrue(panel.txt_coding_preview.isReadOnly())
        self.assertEqual(
            panel.txt_coding_preview.toPlainText(),
            "62 F1 08 37 12",
        )

        panel.table.item(0, 4).setText("FF")
        self.assertEqual(
            panel.txt_coding_preview.toPlainText(),
            "62 F1 08 FF 12",
        )
        self.assertEqual(len(panel.txt_coding_preview.extraSelections()), 1)

        panel.table.item(1, 4).setText("05")
        self.assertEqual(
            panel.txt_coding_preview.toPlainText(),
            "62 F1 08 FF 15",
        )
        self.assertEqual(len(panel.txt_coding_preview.extraSelections()), 1)

    def test_preview_edit_imports_payload_after_check_without_resetting_before(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Editable Byte",
            2,
            0,
            8,
            "0x08=Old\n0x99=New",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("62 F1 08")
            panel.import_coding_value()
            panel.encode_coding_payload()

        self.assertEqual(panel.btn_preview_edit.text(), "EDIT")
        self.assertFalse(panel.btn_preview_edit.isEnabled())
        self.assertTrue(panel.txt_coding_preview.isReadOnly())

        panel.check_coding_value()
        self.assertTrue(panel.btn_preview_edit.isEnabled())

        panel.btn_preview_edit.click()
        self.assertFalse(panel.txt_coding_preview.isReadOnly())
        self.assertEqual(panel.btn_preview_edit.text(), "IMPORT")

        panel.txt_coding_preview.setPlainText("62 F1 99")
        panel.btn_preview_edit.click()

        self.assertTrue(panel.txt_coding_preview.isReadOnly())
        self.assertEqual(panel.btn_preview_edit.text(), "EDIT")
        self.assertEqual(panel.table.item(0, 4).text(), "99")
        self.assertEqual(panel.table.cellWidget(0, 5).currentText(), "New")
        self.assertEqual(panel.table.item(0, 6).text(), "08")
        self.assertEqual(panel.table.item(0, 7).text(), "Old")
        self.assertEqual(panel.table.item(0, 8).text(), "No-M")
        self.assertEqual(panel._payload_baseline_bytes, [0x62, 0xF1, 0x08])
        self.assertEqual(panel._payload_preview_bytes, [0x62, 0xF1, 0x99])

    def test_payload_preview_refresh_copy_and_edit_buttons(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Parameter",
            "BytePos (from 0)",
            "BitPos",
            "BitLength",
            "MethodType",
        ])
        sheet.append([
            "Payload Byte 2",
            2,
            0,
            8,
            "0x08=Old\n0x99=New",
        ])

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "coding.xlsx"
            workbook.save(path)

            panel = CodingValuePanel()
            self.addCleanup(panel.deleteLater)
            panel.file_path.setText(str(path))
            panel.txt_coding_value.setPlainText("62 F1 08")
            panel.import_coding_value()
            panel.encode_coding_payload()

        self.assertEqual(panel.table.item(0, 4).text(), "08")
        panel.txt_coding_value.setPlainText("62 F1 99")

        panel.refresh_coding_preview()
        self.assertEqual(panel.table.item(0, 4).text(), "99")
        self.assertEqual(
            panel.txt_coding_preview.toPlainText(),
            "62 F1 99",
        )

        panel.copy_coding_preview()
        self.assertEqual(
            QApplication.clipboard().text(),
            "62 F1 99",
        )


    def test_shared_byte_editing_preserves_other_nibble(self):
        panel = CodingValuePanel()
        self.addCleanup(panel.deleteLater)
        panel.table.set_rows([
            CodingValueRow(
                parameter="Low Nibble",
                byte_pos="0",
                bit_pos="0",
                bit_length="4",
                raw_value="",
                decoded_value="",
                decoded_options=(
                    CodingValueOption(raw_value="0x05", label="Low Five"),
                    CodingValueOption(raw_value="0x06", label="Low Six"),
                ),
            ),
            CodingValueRow(
                parameter="High Nibble",
                byte_pos="0",
                bit_pos="4",
                bit_length="4",
                raw_value="",
                decoded_value="",
                decoded_options=(
                    CodingValueOption(raw_value="0x0A", label="High A"),
                    CodingValueOption(raw_value="0x0B", label="High B"),
                ),
            ),
        ])
        panel.txt_coding_value.setPlainText("B6")
        panel.encode_coding_payload()

        self.assertEqual(panel.table.item(0, 4).text(), "06")
        self.assertEqual(panel.table.item(1, 4).text(), "0B")
        self.assertEqual(panel.txt_coding_preview.toPlainText(), "B6")

        panel.table.item(0, 4).setText("05")

        self.assertEqual(panel.txt_coding_preview.toPlainText(), "B5")
        self.assertEqual(panel.table.item(1, 4).text(), "0B")

        panel.table.item(1, 4).setText("0A")

        self.assertEqual(panel.txt_coding_preview.toPlainText(), "A5")
        self.assertEqual(panel.table.item(0, 4).text(), "05")

    def test_two_coding_value_panels_keep_instance_state_separate(self):
        panel_a = CodingValuePanel()
        panel_b = CodingValuePanel()
        self.addCleanup(panel_a.deleteLater)
        self.addCleanup(panel_b.deleteLater)

        def row(parameter, byte_pos):
            return CodingValueRow(
                parameter=parameter,
                byte_pos=str(byte_pos),
                bit_pos="0",
                bit_length="8",
                raw_value="",
                decoded_value="",
                decoded_options=(),
            )

        definition_a = CodingDefinition(
            name="Definition A",
            rows=(
                row("Definition A Byte 0", 0),
                row("Definition A Byte 1", 1),
            ),
        )
        definition_b = CodingDefinition(
            name="Definition B",
            rows=(
                row("Definition B Byte 0", 0),
                row("Definition B Byte 1", 1),
            ),
        )
        panel_a._set_coding_definition(definition_a)
        panel_b._set_coding_definition(definition_b)
        panel_a.txt_coding_value.setPlainText("AA 11")
        panel_b.txt_coding_value.setPlainText("BB CC")
        panel_a.encode_coding_payload()
        panel_b.encode_coding_payload()

        self.assertIs(panel_a._coding_definition, definition_a)
        self.assertIs(panel_b._coding_definition, definition_b)
        self.assertIsInstance(panel_a._coding_definition, CodingDefinition)
        self.assertIsInstance(panel_b._coding_definition, CodingDefinition)

        expected_b_input = "BB CC"
        expected_b_preview = "BB CC"
        expected_b_baseline = [0xBB, 0xCC]
        expected_b_raw_values = ["BB", "CC"]
        expected_b_columns = len(CodingValueTable.HEADERS)

        def assert_panel_b_unchanged():
            self.assertIs(panel_b._coding_definition, definition_b)
            self.assertEqual(panel_b.table.rowCount(), 2)
            self.assertEqual(panel_b.table.item(0, 0).text(), "Definition B Byte 0")
            self.assertEqual(panel_b.table.item(1, 0).text(), "Definition B Byte 1")
            self.assertEqual(panel_b.txt_coding_value.toPlainText(), expected_b_input)
            self.assertEqual(panel_b.txt_coding_preview.toPlainText(), expected_b_preview)
            self.assertEqual(panel_b._payload_baseline_bytes, expected_b_baseline)
            self.assertEqual(panel_b._payload_preview_bytes, expected_b_baseline)
            self.assertEqual(
                [panel_b.table.item(row_index, 4).text() for row_index in range(2)],
                expected_b_raw_values,
            )
            self.assertEqual(panel_b.table.columnCount(), expected_b_columns)
            self.assertFalse(panel_b.table.isRowHidden(0))
            self.assertFalse(panel_b.table.isRowHidden(1))
            self.assertEqual(panel_b.txt_parameter_filter.text(), "")

        self.assertEqual(panel_a.table.item(0, 0).text(), "Definition A Byte 0")
        self.assertEqual(panel_a.table.item(1, 0).text(), "Definition A Byte 1")
        self.assertEqual(panel_a._payload_baseline_bytes, [0xAA, 0x11])
        self.assertEqual(panel_a._payload_preview_bytes, [0xAA, 0x11])
        assert_panel_b_unchanged()

        panel_a.table.fn_set_raw_value_at_row(0, "AB")
        self.assertEqual(panel_a.txt_coding_preview.toPlainText(), "AB 11")
        self.assertEqual(panel_a._payload_preview_bytes, [0xAB, 0x11])
        assert_panel_b_unchanged()

        panel_a.check_coding_value()
        self.assertGreater(panel_a.table.columnCount(), len(CodingValueTable.HEADERS))
        assert_panel_b_unchanged()

        panel_a.filter_no_match_rows()
        self.assertEqual(panel_a.txt_parameter_filter.text(), "No-M")
        assert_panel_b_unchanged()

        panel_a.clear_coding_payload()
        self.assertEqual(panel_a.txt_coding_value.toPlainText(), "")
        self.assertEqual(panel_a.txt_coding_preview.toPlainText(), "")
        assert_panel_b_unchanged()

        replacement_definition_a = CodingDefinition(
            name="Definition A Replacement",
            rows=(row("Replacement A Byte 0", 0),),
        )
        panel_a._set_coding_definition(replacement_definition_a)
        self.assertIs(panel_a._coding_definition, replacement_definition_a)
        self.assertEqual(panel_a.table.item(0, 0).text(), "Replacement A Byte 0")
        assert_panel_b_unchanged()
if __name__ == "__main__":
    unittest.main()
