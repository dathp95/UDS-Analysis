import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QDialog

from gui.dialogs.qr_list_dialog import QRListDialog
from gui.tabs.crc_converter_tab import CRCConverterTab
from gui.themes.theme import ThemeType
from gui.themes.theme_manager import ThemeManager


class QRListDialogTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def dialog(self):
        dialog = QRListDialog()
        self.addCleanup(dialog.deleteLater)
        return dialog

    def test_count_limits_and_warning_without_truncation(self):
        dialog = self.dialog()
        self.assertEqual(dialog.windowTitle(), "Import List QR Code")
        self.assertEqual(dialog.lbl_count.text(), "Total QR Codes: 0 / 50")
        self.assertFalse(dialog.btn_create.isEnabled())
        for count in (1, 50, 51, 50, 0):
            dialog.txt_codes.setPlainText("\n".join(["00001"] * count))
            self.assertEqual(dialog.lbl_count.text(), f"Total QR Codes: {count} / 50")
            self.assertEqual(dialog.btn_create.isEnabled(), 1 <= count <= 50)
            self.assertEqual(not dialog.lbl_limit_warning.isHidden(), count > 50)
            self.assertEqual(len(dialog.txt_codes.toPlainText().splitlines()), count)

    def test_create_real_workbook_and_report_sheet_count(self):
        dialog = self.dialog()
        dialog.txt_codes.setPlainText("\n".join(f"QR-{i:03}" for i in range(50)))
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[2]) as directory:
            output = Path(directory) / "list.xlsx"
            with patch("gui.dialogs.qr_list_dialog.QFileDialog.getSaveFileName", return_value=(str(output), "")), patch(
                "gui.dialogs.qr_list_dialog.QMessageBox.information"
            ) as message:
                dialog.btn_create.click()
            self.assertTrue(output.exists())
            self.assertIn("Total QR Codes: 50", message.call_args.args[2])
            self.assertIn("Total Sheets: 3", message.call_args.args[2])
            book = load_workbook(output)
            try:
                self.assertEqual(book.sheetnames, ["QR_001", "QR_002", "QR_003"])
            finally:
                book.close()
        self.assertTrue(dialog.btn_create.isEnabled())
        self.assertTrue(dialog.btn_cancel.isEnabled())

    def test_cancel_save_dialog_has_no_export_side_effect(self):
        dialog = self.dialog()
        dialog.txt_codes.setPlainText("QR")
        with patch("gui.dialogs.qr_list_dialog.QFileDialog.getSaveFileName", return_value=("", "")), patch(
            "gui.dialogs.qr_list_dialog.export_qr_codes_to_excel"
        ) as export:
            dialog.btn_create.click()
        export.assert_not_called()
        self.assertTrue(dialog.btn_create.isEnabled())

    def test_default_filename_and_missing_extension(self):
        dialog = self.dialog()
        dialog.txt_codes.setPlainText("00001")
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[2]) as directory:
            output = Path(directory) / "export"
            with patch("gui.dialogs.qr_list_dialog.QFileDialog.getSaveFileName", return_value=(str(output), "")) as choose, patch(
                "gui.dialogs.qr_list_dialog.QMessageBox.information"
            ):
                dialog.btn_create.click()
            self.assertRegex(choose.call_args.args[2], r"^QR_Code_List_\d{8}_\d{6}\.xlsx$")
            self.assertTrue(output.with_suffix(".xlsx").exists())

    def test_cancel_closes_dialog_without_export(self):
        dialog = self.dialog()
        with patch("gui.dialogs.qr_list_dialog.export_qr_codes_to_excel") as export:
            dialog.btn_cancel.click()
        export.assert_not_called()
        self.assertEqual(dialog.result(), QDialog.DialogCode.Rejected)

    def test_invalid_count_cannot_bypass_validation_or_open_save_dialog(self):
        dialog = self.dialog()
        for count in (0, 51):
            dialog.txt_codes.setPlainText("\n".join(["QR"] * count))
            with patch("gui.dialogs.qr_list_dialog.QFileDialog.getSaveFileName") as choose, patch(
                "gui.dialogs.qr_list_dialog.QMessageBox.warning"
            ) as warning:
                dialog.create_excel()
            choose.assert_not_called()
            warning.assert_called_once()

    def test_export_error_restores_controls_and_prevents_reentry(self):
        dialog = self.dialog()
        dialog.txt_codes.setPlainText("QR")

        def fail(*_args):
            self.assertFalse(dialog.btn_create.isEnabled())
            self.assertFalse(dialog.txt_codes.isEnabled())
            dialog.create_excel()
            raise OSError("Access denied")

        with patch("gui.dialogs.qr_list_dialog.QFileDialog.getSaveFileName", return_value=("example.xlsx", "")), patch(
            "gui.dialogs.qr_list_dialog.export_qr_codes_to_excel", side_effect=fail
        ) as export, patch("gui.dialogs.qr_list_dialog.QMessageBox.warning") as warning:
            dialog.btn_create.click()
        export.assert_called_once()
        warning.assert_called_once()
        self.assertTrue(dialog.btn_create.isEnabled())
        self.assertTrue(dialog.btn_cancel.isEnabled())
        self.assertTrue(dialog.txt_codes.isEnabled())

    def test_light_dark_theme_preserves_data_and_limit_state(self):
        dialog = self.dialog()
        dialog.txt_codes.setPlainText("\n".join(["QR"] * 51))
        before = dialog.txt_codes.toPlainText()
        self.addCleanup(ThemeManager.fn_set_theme, ThemeManager.fn_theme())
        for theme in (ThemeType.DARK, ThemeType.LIGHT):
            ThemeManager.fn_set_theme(theme)
            dialog.fn_refresh_theme()
            self.assertIn(ThemeManager.fn_colors().WINDOW, dialog.txt_codes.styleSheet())
            self.assertIn(ThemeManager.fn_colors().DANGER, dialog.lbl_limit_warning.styleSheet())
            self.assertEqual(dialog.txt_codes.toPlainText(), before)
            self.assertFalse(dialog.btn_create.isEnabled())

    def test_crc_integration_button_opens_dialog_without_changing_single_qr(self):
        tab = CRCConverterTab()
        self.addCleanup(tab.deleteLater)
        tab.txt_qr_input.setPlainText("Single QR stays unchanged")
        self.assertEqual(tab.btn_import_qr_list.text(), "Import List QR Code")
        layout = tab.btn_clear_qr.parentWidget().findChildren(type(tab.layout()))
        self.assertTrue(any(row.indexOf(tab.btn_import_qr_list) == row.indexOf(tab.btn_clear_qr) + 1
                            for row in layout if row.indexOf(tab.btn_clear_qr) >= 0))
        with patch("gui.tabs.crc_converter_tab.QRListDialog") as dialog:
            tab.btn_import_qr_list.click()
        dialog.assert_called_once_with(tab)
        dialog.return_value.exec.assert_called_once()
        dialog.return_value.deleteLater.assert_called_once()
        self.assertEqual(tab.txt_qr_input.toPlainText(), "Single QR stays unchanged")
