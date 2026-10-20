import csv
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from core.build_transactions import build_transactions_v4
from core.report_engine import build_ecu_reports, build_summary_report
from core.report_export import export_workbook
from core.transaction_formatter import fn_build_transaction_rows
from gui.controllers.clipboard_controller import ClipboardController
from gui.tabs.log_analyzer_tab import LogAnalyzerTab
from gui.themes.theme import ThemeType
from gui.themes.theme_manager import ThemeManager
from gui.widgets.log_analyzer.result_table import ResultTable
from gui.widgets.log_analyzer.action_panel import ActionPanel
from tests.core.test_transaction_matching import msg


class CopyDataTableTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.application = QApplication.instance() or QApplication([])

    def setUp(self):
        self.table = ResultTable()
        self.addCleanup(self.table.deleteLater)
        self.controller = ClipboardController()
        self.transactions = build_transactions_v4(
            [msg(1.234567, 0x682, "22 F1 90"), msg(2.123456, 0x693, "22 F1 91"),
             msg(4, 0x682, "31 01 02 06")],
            [msg(1.246912, 0x602, "62 F1 90 " + "AB " * 512)],
            [msg(2.223456, 0x613, "7F 22 31"), msg(4.1, 0x602, "7F 31 78")],
            {"A": {"request": 0x682, "response": 0x602},
             "B": {"request": 0x693, "response": 0x613}}, 10)
        self.transactions[0]["display_name"] = 'Read "first"'
        self.rows = fn_build_transaction_rows(self.transactions)
        self.table.set_data(self.rows)
        self.original_clipboard = self.application.clipboard().text()
        self.addCleanup(self.application.clipboard().setText, self.original_clipboard)

    def copy(self):
        with patch("gui.controllers.clipboard_controller.QMessageBox.information") as info, \
             patch("gui.controllers.clipboard_controller.QMessageBox.warning") as warning:
            self.controller.fn_copy(parent=self.table, rows=self.table.fn_visible_rows())
        warning.assert_not_called()
        return self.application.clipboard().text(), info

    def parsed(self, text):
        return list(csv.reader(io.StringIO(text), dialect="excel-tab"))

    def test_all_rows_match_actual_all_ecus_sheet_and_ignore_selection(self):
        self.table.selectRow(1)
        text, info = self.copy()
        copied = self.parsed(text)
        self.assertEqual(len(copied), 4)
        self.assertIn("3", info.call_args.args[2])
        self.assertEqual(copied[0], list(self.rows[0]))
        self.assertEqual(copied[1][-2], str(self.rows[0]["RT (ms)"]))
        self.assertEqual(copied[1][4], self.rows[0]["Response"])
        self.assertEqual(copied[3][4:6], ["", ""])
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[2]) as directory:
            path = Path(directory) / "expected.xlsx"
            reports = build_ecu_reports(self.transactions)
            self.assertTrue(export_workbook(build_summary_report(reports), reports, self.transactions, path))
            workbook = load_workbook(path)
            try:
                expected = list(workbook["ALL_ECUS"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(len(copied), len(expected))
                for actual_row, expected_row in zip(copied, expected):
                    for actual, value in zip(actual_row, expected_row):
                        if isinstance(value, (int, float)):
                            self.assertEqual(float(actual), value)
                        else:
                            self.assertEqual(actual, "" if value is None else value)
            finally:
                workbook.close()

    def test_quick_filter_copies_only_visible_transactions(self):
        self.table.fn_apply_quick_filter({"filters": {"request": "22 F1 90"}})
        copied = self.parsed(self.copy()[0])
        self.assertEqual(len(copied), 2)
        self.assertEqual(copied[1][3], "22 F1 90")

    def test_search_copies_only_visible_transactions(self):
        self.table.fn_search("PENDING_TIMEOUT")
        copied = self.parsed(self.copy()[0])
        self.assertEqual(len(copied), 2)
        self.assertEqual(copied[1][-1], "PENDING_TIMEOUT")

    def test_ecu_filter_preserves_display_order_and_ignores_hidden_selection(self):
        self.table.fn_apply_quick_filter({"filters": {"ecu": "A"}})
        self.table.selectRow(1)
        copied = self.parsed(self.copy()[0])
        self.assertEqual([row[3] for row in copied[1:]], ["22 F1 90", "31 01 02 06"])

    def test_empty_filter_result_and_empty_table_keep_clipboard(self):
        for clear in (False, True):
            with self.subTest(clear=clear):
                if clear:
                    self.table.clear_data()
                else:
                    self.table.fn_search("no matching transaction")
                self.application.clipboard().setText("keep existing clipboard")
                text, info = self.copy()
                self.assertEqual(text, "keep existing clipboard")
                self.assertIn("No visible transactions", info.call_args.args[2])

    def test_control_characters_quotes_unicode_and_missing_cells(self):
        self.rows[0]["Activity"] = 'one\ttwo\r\nthree\nfour\rfive "quoted" \u0111'
        self.rows[1]["Response"] = None
        self.table.set_data(self.rows)
        self.table.takeItem(1, 2)
        text, _ = self.copy()
        copied = self.parsed(text)
        self.assertEqual(len(text.splitlines()), 4)
        self.assertTrue(all(len(row) == 7 for row in copied))
        self.assertEqual(copied[1][2], 'one two three four five "quoted" \u0111')
        self.assertEqual(copied[2][4], "")
        self.assertEqual(copied[2][2], "")

    def test_light_and_dark_copy_identical_text(self):
        self.addCleanup(ThemeManager.fn_set_theme, ThemeManager.fn_theme())
        outputs = []
        for theme in (ThemeType.LIGHT, ThemeType.DARK):
            ThemeManager.fn_set_theme(theme)
            self.table.fn_refresh_theme()
            outputs.append(self.copy()[0])
        self.assertEqual(outputs[0], outputs[1])

    def test_copy_button_renders_full_label_in_both_themes(self):
        self.addCleanup(ThemeManager.fn_set_theme, ThemeManager.fn_theme())
        panel = ActionPanel()
        self.addCleanup(panel.deleteLater)
        panel.resize(220, 800)
        panel.show()
        for theme in (ThemeType.LIGHT, ThemeType.DARK):
            with self.subTest(theme=theme):
                ThemeManager.fn_set_theme(theme)
                panel.fn_refresh_theme()
                self.application.processEvents()
                button = panel.btn_copy
                self.assertEqual(button.text(), "COPY DATA TABLE")
                self.assertLessEqual(button.fontMetrics().horizontalAdvance(button.text()) + 24,
                                     button.width())
                self.assertFalse(button.grab().isNull())

    def test_tsv_preserves_zero_none_missing_values_and_one_header(self):
        from core.transaction_formatter import TRANSACTION_COLUMNS, fn_format_transaction_rows_tsv
        copied = self.parsed(fn_format_transaction_rows_tsv([
            {"ECU": "A", "Time": 0, "RT (ms)": 0, "Response": None}]))
        self.assertEqual(copied[0], list(TRANSACTION_COLUMNS))
        self.assertEqual(copied[1], ["A", "0", "", "", "", "0", ""])

    def test_clipboard_error_shows_warning_without_success_notification(self):
        with patch("gui.controllers.clipboard_controller.QApplication.clipboard",
                   side_effect=RuntimeError("Clipboard unavailable")), \
             patch("gui.controllers.clipboard_controller.QMessageBox.warning") as warning, \
             patch("gui.controllers.clipboard_controller.QMessageBox.information") as info:
            self.controller.fn_copy(parent=self.table, rows=self.table.fn_visible_rows())
        info.assert_not_called()
        self.assertEqual(warning.call_args.args[2], "Clipboard unavailable")

    def test_reordered_display_rows_are_copied_in_current_order(self):
        self.table.set_data(list(reversed(self.rows)))
        copied = self.parsed(self.copy()[0])
        self.assertEqual([row[3] for row in copied[1:]],
                         [row["Request"] for row in reversed(self.rows)])

    def test_copy_button_signal_and_no_analysis_lifecycle(self):
        tab = LogAnalyzerTab()
        self.addCleanup(tab.deleteLater)
        panel = tab.right_panel.action_panel
        self.assertEqual(panel.btn_copy.text(), "COPY DATA TABLE")
        self.assertFalse(panel.btn_copy.isEnabled())
        tab._fn_analysis_finished({"transactions": []})
        self.assertFalse(panel.btn_copy.isEnabled())
        tab._fn_analysis_finished({"transactions": self.transactions})
        self.assertTrue(panel.btn_copy.isEnabled())
        tab.tbl_result.fn_search("PENDING_TIMEOUT")
        with patch("gui.controllers.clipboard_controller.QMessageBox.information"):
            panel.btn_copy.click()
        self.assertEqual(self.parsed(self.application.clipboard().text())[1][-1], "PENDING_TIMEOUT")
        tab.fn_clear_clicked()
        self.assertFalse(panel.btn_copy.isEnabled())


if __name__ == "__main__":
    unittest.main()
