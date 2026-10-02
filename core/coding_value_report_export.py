from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook

from core.excel_styles import (
    HEADER_ALIGNMENT,
    HEADER_FILL,
    HEADER_FONT,
    THIN_BORDER,
    WARNING_FILL,
)


CODING_VALUE_REPORT_HEADERS = [
    "Parameter",
    "Byte Pos",
    "Bit Pos",
    "Bit Lengh",
    "Raw value",
    "Decoded Value (Editable)",
    "Raw value (before)",
    "Decoded value (before)",
    "Result",
]


@dataclass(frozen=True)
class CodingValueReportRow:
    parameter: str
    byte_pos: str
    bit_pos: str
    bit_length: str
    raw_value: str
    decoded_value: str
    raw_value_before: str = ""
    decoded_value_before: str = ""
    result: str = ""

    def values_for_headers(self, headers: Sequence[str]) -> list[str]:
        values_by_header = {
            "Parameter": self.parameter,
            "Byte Pos": self.byte_pos,
            "Bit Pos": self.bit_pos,
            "Bit Lengh": self.bit_length,
            "Raw value": self.raw_value,
            "Decoded Value (Editable)": self.decoded_value,
            "Raw value (before)": self.raw_value_before,
            "Decoded value (before)": self.decoded_value_before,
            "Result": self.result,
        }
        return [values_by_header.get(header, "") for header in headers]


@dataclass(frozen=True)
class CodingValueReportData:
    sheet_name: str
    headers: Sequence[str]
    rows: Sequence[CodingValueReportRow]


def export_coding_value_report(
        report_data: CodingValueReportData,
        output_file: str | Path,
    ) -> Path:
    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = safe_excel_sheet_name(report_data.sheet_name)
    headers = list(report_data.headers)
    sheet.append(headers)
    result_index = _result_column_index(headers)
    no_match_rows = []

    for row in report_data.rows:
        values = row.values_for_headers(headers)
        sheet.append(values)
        if _is_no_match_export_row(values, result_index):
            no_match_rows.append(sheet.max_row)

    _format_coding_export_sheet(sheet)

    for row_number in no_match_rows:
        _apply_no_match_export_fill(sheet[row_number])

    workbook.save(output_path)
    return output_path


def safe_excel_sheet_name(name: str) -> str:
    safe_name = str(name or "")
    invalid_chars = r'[]:*?/\\'
    for char in invalid_chars:
        safe_name = safe_name.replace(char, "_")

    return safe_name[:31] or "Coding value"


def _result_column_index(headers: Sequence[str]) -> int | None:
    try:
        return list(headers).index("Result")
    except ValueError:
        return None


def _is_no_match_export_row(values: Sequence[str], result_index: int | None) -> bool:
    return (
        result_index is not None
        and result_index < len(values)
        and values[result_index] == "No-M"
    )


def _format_coding_export_sheet(sheet) -> None:
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER

    for row in sheet.iter_rows(
        min_row=2,
        max_row=sheet.max_row,
        min_col=1,
        max_col=sheet.max_column,
    ):
        for cell in row:
            cell.border = THIN_BORDER

    sheet.freeze_panes = "A2"

    if sheet.max_column > 0:
        sheet.auto_filter.ref = sheet.dimensions


def _apply_no_match_export_fill(row_cells) -> None:
    for cell in row_cells:
        cell.fill = WARNING_FILL
