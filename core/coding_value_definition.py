from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from core.coding_value import CodingDefinition, CodingValueOption, CodingValueRow
from core.coding_value_payload import extract_raw_value, hex_tokens, parse_payload_text, to_int


_OPTION_PATTERN = re.compile(
    r"(?P<raw>0[xX][0-9A-Fa-f]+)\s*[=:-]\s*(?P<label>[^\n\r]+)"
)


def load_coding_definition_from_excel(
        excel_path: str | Path,
        coding_payload: str = "",
    ) -> CodingDefinition:
    source_path = Path(excel_path)
    workbook = load_workbook(
        source_path,
        data_only=True,
        read_only=True,
    )
    payload_bytes = parse_payload_text(coding_payload)
    rows: list[CodingValueRow] = []

    try:
        for sheet in workbook.worksheets:
            sheet_values = [
                [cell for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            rows.extend(
                _load_parameter_table_rows(
                    sheet_values,
                    payload_bytes,
                )
            )
            if not rows:
                rows.extend(
                    _load_dollar_header_rows(
                        sheet_values,
                        payload_bytes,
                    )
                )
    finally:
        workbook.close()

    return CodingDefinition(
        name=source_path.stem or "coding_value",
        rows=tuple(rows),
        source_file=source_path.name,
        source_path=str(source_path),
    )


def load_coding_definition_from_json(json_path: str | Path) -> CodingDefinition:
    path = Path(json_path)
    payload = json.loads(
        path.read_text(encoding="utf-8")
    )
    source_file = str(payload.get("source_file") or "")
    source_path = str(payload.get("source_path") or "")
    rows = tuple(
        _coding_value_row_from_json(row)
        for row in payload.get("rows", [])
    )

    return CodingDefinition(
        name=_definition_name_from_json(path, source_file),
        rows=rows,
        source_file=source_file,
        source_path=source_path,
    )


def load_coding_definition(
        path: str | Path,
        coding_payload: str = "",
    ) -> CodingDefinition:
    source_path = Path(path)
    if source_path.suffix.lower() == ".json":
        return load_coding_definition_from_json(source_path)

    return load_coding_definition_from_excel(
        source_path,
        coding_payload,
    )


def export_coding_definition_to_json(
        definition: CodingDefinition,
        output_path: str | Path,
    ) -> Path:
    output = Path(output_path)
    if output.suffix.lower() == ".json":
        json_path = output
    else:
        json_path = output / f"{_definition_export_stem(definition)}.json"

    json_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )
    payload = {
        "source_file": definition.source_file,
        "source_path": definition.source_path,
        "rows": [
            _coding_value_row_to_json(row)
            for row in definition.rows
        ],
    }
    json_path.write_text(
        json.dumps(payload, indent=4, ensure_ascii=False),
        encoding="utf-8",
    )
    return json_path


def _definition_export_stem(definition: CodingDefinition) -> str:
    if definition.source_file:
        stem = Path(definition.source_file).stem
        if stem:
            return stem

    return definition.name or "coding_value"


def _definition_name_from_json(json_path: Path, source_file: str) -> str:
    if source_file:
        stem = Path(source_file).stem
        if stem:
            return stem

    return json_path.stem or "coding_value"


def _coding_value_row_from_json(row: dict[str, Any]) -> CodingValueRow:
    return CodingValueRow(
        parameter=str(row.get("parameter", "")),
        byte_pos=str(row.get("byte_pos", "")),
        bit_pos=str(row.get("bit_pos", "")),
        bit_length=str(row.get("bit_length", "")),
        raw_value=str(row.get("raw_value", "")),
        decoded_value=str(row.get("decoded_value", "")),
        decoded_options=tuple(
            CodingValueOption(
                raw_value=str(option.get("raw_value", "")),
                label=str(option.get("label", "")),
            )
            for option in row.get("decoded_options", [])
        ),
    )


def _coding_value_row_to_json(row: CodingValueRow) -> dict[str, Any]:
    return {
        "parameter": row.parameter,
        "byte_pos": row.byte_pos,
        "bit_pos": row.bit_pos,
        "bit_length": row.bit_length,
        "raw_value": row.raw_value,
        "decoded_value": row.decoded_value,
        "decoded_options": [
            {
                "raw_value": option.raw_value,
                "label": option.label,
            }
            for option in row.decoded_options
        ],
    }


def _load_parameter_table_rows(
        sheet_values: list[list[Any]],
        payload_bytes: list[int],
    ) -> list[CodingValueRow]:
    header_index = _find_parameter_header_row(sheet_values)
    if header_index is None:
        return []

    headers = _parameter_table_headers(
        sheet_values[header_index]
    )
    parameter_col = headers.get("parameter")
    byte_col = headers.get("bytepos") or headers.get("byteposfrom0")
    bit_col = headers.get("bitpos")
    length_col = headers.get("bitlength")
    method_col = headers.get("methodtype")

    if (
            parameter_col is None
            or byte_col is None
            or bit_col is None
            or length_col is None
        ):
        return []

    rows: list[CodingValueRow] = []
    for row in sheet_values[header_index + 1:]:
        parameter = _cell_text(_get_cell(row, parameter_col)).strip()
        if not parameter:
            continue

        rows.append(
            _build_row(
                parameter=parameter.lstrip("$").strip(),
                byte_pos=_get_cell(row, byte_col),
                bit_pos=_get_cell(row, bit_col),
                bit_length=_get_cell(row, length_col),
                method_type=_get_cell(row, method_col) if method_col is not None else "",
                payload_bytes=payload_bytes,
            )
        )

    return rows


def _load_dollar_header_rows(
        sheet_values: list[list[Any]],
        payload_bytes: list[int],
    ) -> list[CodingValueRow]:
    if not sheet_values:
        return []

    header_row = sheet_values[0]
    rows: list[CodingValueRow] = []
    for column, header in enumerate(header_row):
        if not isinstance(header, str) or not header.strip().startswith("$"):
            continue

        values_by_key = {
            _normalize_header(_get_cell(row, 0)): _get_cell(row, column)
            for row in sheet_values[1:]
        }
        rows.append(
            _build_row(
                parameter=header.strip().lstrip("$").strip(),
                byte_pos=values_by_key.get("bytepos", ""),
                bit_pos=values_by_key.get("bitpos", ""),
                bit_length=(
                    values_by_key.get("bitlength", "")
                    or values_by_key.get("bitlengh", "")
                ),
                method_type=(
                    values_by_key.get("methodtype", "")
                    or values_by_key.get("decodedvalue", "")
                ),
                payload_bytes=payload_bytes,
            )
        )

    return rows


def _build_row(
        parameter: str,
        byte_pos: Any,
        bit_pos: Any,
        bit_length: Any,
        method_type: Any,
        payload_bytes: list[int],
    ) -> CodingValueRow:
    raw_value = _extract_definition_raw_value(
        payload_bytes,
        to_int(byte_pos),
        to_int(bit_pos),
        to_int(bit_length),
    )
    options = _parse_options(method_type)
    decoded_value = _decode_raw_value(
        raw_value,
        options,
    )

    return CodingValueRow(
        parameter=parameter,
        byte_pos=_cell_text(byte_pos),
        bit_pos=_cell_text(bit_pos),
        bit_length=_cell_text(bit_length),
        raw_value=raw_value,
        decoded_value=decoded_value,
        decoded_options=options,
    )


def _parameter_table_headers(header_row: list[Any]) -> dict[str, int]:
    if _has_required_dollar_headers(header_row):
        return {
            _normalize_header(str(value).strip().lstrip("$")): index
            for index, value in enumerate(header_row)
            if isinstance(value, str) and value.strip().startswith("$")
        }

    return {
        _normalize_header(value): index
        for index, value in enumerate(header_row)
        if value is not None
    }


def _has_required_dollar_headers(header_row: list[Any]) -> bool:
    dollar_headers = {
        _normalize_header(str(value).strip().lstrip("$"))
        for value in header_row
        if isinstance(value, str) and value.strip().startswith("$")
    }
    return (
        "parameter" in dollar_headers
        and (
            "bytepos" in dollar_headers
            or "byteposfrom0" in dollar_headers
        )
        and "bitpos" in dollar_headers
        and "bitlength" in dollar_headers
    )


def _find_parameter_header_row(sheet_values: list[list[Any]]) -> int | None:
    for index, row in enumerate(sheet_values):
        normalized = {
            _normalize_header(value)
            for value in row
            if value is not None
        }
        if (
                "parameter" in normalized
                and (
                    "bytepos" in normalized
                    or "byteposfrom0" in normalized
                )
                and "bitpos" in normalized
                and "bitlength" in normalized
            ):
            return index

    return None


def _parse_options(value: Any) -> tuple[CodingValueOption, ...]:
    text = _cell_text(value)
    options = []
    for match in _OPTION_PATTERN.finditer(text):
        options.append(
            CodingValueOption(
                raw_value=_normalize_raw_hex(match.group("raw")),
                label=match.group("label").strip(),
            )
        )

    return tuple(options)


def _decode_raw_value(
        raw_value: str,
        options: tuple[CodingValueOption, ...],
    ) -> str:
    if not raw_value:
        return ""

    raw_int = _hex_to_int(raw_value)
    for option in options:
        if _hex_to_int(option.raw_value) == raw_int:
            return option.label

    return raw_value


def _extract_definition_raw_value(
        payload_bytes: list[int],
        byte_pos: int | None,
        bit_pos: int | None,
        bit_length: int | None,
    ) -> str:
    raw_value = extract_raw_value(
        payload_bytes,
        byte_pos,
        bit_pos,
        bit_length,
    )
    if not raw_value:
        return ""

    tokens = hex_tokens(raw_value)
    if not tokens:
        return ""

    raw_int = int("".join(tokens), 16)
    if bit_length is not None and bit_length > 0:
        width = max(1, (bit_length + 3) // 4)
    else:
        width = max(1, len("".join(tokens)))

    return f"0x{raw_int:0{width}X}"


def _normalize_header(value: Any) -> str:
    return re.sub(
        r"[^0-9a-z]+",
        "",
        _cell_text(value).lower(),
    )


def _normalize_raw_hex(value: str) -> str:
    raw_int = _hex_to_int(value)
    if raw_int is None:
        return value.strip()

    width = max(1, len(value.replace("0x", "").replace("0X", "")))
    return f"0x{raw_int:0{width}X}"


def _hex_to_int(value: str) -> int | None:
    try:
        return int(value, 16)
    except (TypeError, ValueError):
        return None


def _get_cell(row: list[Any], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None

    return row[index]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()
