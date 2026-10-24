"""Synchronous, bounded QR workbook export without GUI dependencies."""
from collections.abc import Iterable
from contextlib import ExitStack
from dataclasses import dataclass
from io import BytesIO
from math import ceil
import os
from pathlib import Path
import re
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorkbookImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.page import PageMargins

from core.qr_generator import generate_qr_png_bytes

MAX_QR_CODES = 50
QR_CODES_PER_SHEET = 24
LIMIT_MESSAGE = "Maximum 50 QR codes per export.\nPlease reduce the list before creating Excel."
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


@dataclass(frozen=True)
class QRExcelExportResult:
    output_file: Path
    code_count: int
    sheet_count: int


def parse_qr_codes(text: str) -> list[str]:
    """Trim lines without deduplication or changes to internal whitespace."""
    return [value for line in text.splitlines() if (value := line.strip())]


def _validate_codes(codes: Iterable[str] | str) -> list[str]:
    values = []
    for code in (parse_qr_codes(codes) if isinstance(codes, str) else codes):
        if not isinstance(code, str):
            raise ValueError("QR code values must be strings.")
        value = code.strip()
        if not value:
            continue
        values.append(value)
        if len(values) > MAX_QR_CODES:
            raise ValueError(LIMIT_MESSAGE)
        if len(value) > 32767 or _INVALID_XML.search(value):
            raise ValueError("QR code text cannot be represented safely in an Excel cell.")
    if not values:
        raise ValueError("Enter at least one QR code.")
    return values


def _configure_sheet(sheet):
    sheet.sheet_view.showGridLines = False
    for column in ("A", "B", "C"):
        sheet.column_dimensions[column].width = 28
    sheet.page_setup.orientation = sheet.ORIENTATION_PORTRAIT
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.25, bottom=0.25,
                                     header=0, footer=0)
    sheet.print_options.horizontalCentered = True


def export_qr_codes_to_excel(codes: Iterable[str] | str, output_file: str | Path) -> QRExcelExportResult:
    """Export 1-50 codes atomically. Raise validation or I/O errors on failure."""
    values = _validate_codes(codes)
    destination = Path(output_file).expanduser().absolute()
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("Select an .xlsx output file.")
    if not destination.parent.is_dir():
        raise FileNotFoundError("The output directory does not exist.")
    if destination.is_dir():
        raise IsADirectoryError("The output path is a directory.")

    temporary_path = None
    try:
        with ExitStack() as resources:
            workbook = Workbook()
            resources.callback(workbook.close)
            workbook.remove(workbook.active)
            for index, value in enumerate(values):
                position = index % QR_CODES_PER_SHEET
                if position == 0:
                    sheet = workbook.create_sheet(f"QR_{index // QR_CODES_PER_SHEET + 1:03d}")
                    _configure_sheet(sheet)
                row, column = (position // 3) * 2 + 1, position % 3 + 1
                sheet.row_dimensions[row].height = 108
                # Leave room for wrapped labels without changing the stored value.
                label_height = min(409, max(24, ceil(len(value) / 26) * 12 + 6))
                sheet.row_dimensions[row + 1].height = max(
                    sheet.row_dimensions[row + 1].height or 0, label_height)
                buffer = resources.enter_context(BytesIO(generate_qr_png_bytes(value)))
                image = WorkbookImage(buffer)
                image.width = image.height = 135
                image.anchor = OneCellAnchor(
                    _from=AnchorMarker(col=column - 1, row=row - 1,
                                       colOff=pixels_to_EMU(33), rowOff=pixels_to_EMU(4)),
                    ext=XDRPositiveSize2D(pixels_to_EMU(135), pixels_to_EMU(135)),
                )
                sheet.add_image(image)
                label = sheet.cell(row + 1, column, value)
                # Explicit string storage also protects values starting with '=' or '#N/A'.
                label.data_type = "s"
                label.number_format = "@"
                label.quotePrefix = value.startswith(("=", "+", "-", "@"))
                label.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                label.font = Font(name="Arial", size=10)
                sheet.print_area = f"A1:C{row + 1}"

            with NamedTemporaryFile(dir=destination.parent, prefix=f".{destination.name}.",
                                    suffix=".tmp.xlsx", delete=False) as temporary:
                temporary_path = Path(temporary.name)
            # Owning the stream ensures Windows can remove a failed partial save.
            with temporary_path.open("wb") as stream:
                workbook.save(stream)
            os.replace(temporary_path, destination)
        return QRExcelExportResult(destination, len(values), ceil(len(values) / QR_CODES_PER_SHEET))
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
