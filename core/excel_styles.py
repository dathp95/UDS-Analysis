from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)


# ==========================================================
# Colors
# ==========================================================

COLOR_HEADER = "D9EAF7"
COLOR_SECTION = "1F4E78"
COLOR_BORDER = "B7B7B7"
COLOR_PASS = "C6EFCE"
COLOR_FAIL = "FFC7CE"
COLOR_WARNING = "E67E22"
COLOR_WHITE = "FFFFFF"


# ==========================================================
# Fills
# ==========================================================

HEADER_FILL = PatternFill(
    "solid",
    fgColor=COLOR_HEADER,
)

SECTION_FILL = PatternFill(
    "solid",
    fgColor=COLOR_SECTION,
)

PASS_FILL = PatternFill(
    "solid",
    fgColor=COLOR_PASS,
)

FAIL_FILL = PatternFill(
    "solid",
    fgColor=COLOR_FAIL,
)

WARNING_FILL = PatternFill(
    "solid",
    fgColor=COLOR_WARNING,
)


# ==========================================================
# Borders
# ==========================================================

THIN_BORDER = Border(
    left=Side(
        style="thin",
        color=COLOR_BORDER,
    ),
    right=Side(
        style="thin",
        color=COLOR_BORDER,
    ),
    top=Side(
        style="thin",
        color=COLOR_BORDER,
    ),
    bottom=Side(
        style="thin",
        color=COLOR_BORDER,
    ),
)


# ==========================================================
# Fonts
# ==========================================================

HEADER_FONT = Font(
    bold=True,
)

SECTION_FONT = Font(
    bold=True,
    color=COLOR_WHITE,
)


# ==========================================================
# Alignment
# ==========================================================

HEADER_ALIGNMENT = Alignment(
    horizontal="center",
    vertical="center",
)

CENTER_ALIGNMENT = Alignment(
    horizontal="center",
    vertical="center",
)

LEFT_ALIGNMENT = Alignment(
    horizontal="left",
    vertical="center",
)
